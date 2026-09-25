"""PPT parser.
Contains parsers for presentation (.pptx) files to extract slide text.
"""
import re
from pathlib import Path
from typing import Any, Dict, List, Union

from docsgpt.parser.file.base_parser import BaseParser

# A GFM row ends at a line break and its columns are separated by pipes, so
# neither can survive inside a cell.
_CELL_BREAK = re.compile(r"\s*(?:\r\n|\r|\n)\s*")
_CELL_PIPE = re.compile(r"(?<!\\)(\\*)\|")


class PPTXParser(BaseParser):
    r"""PPTX (.pptx) parser for extracting text from PowerPoint slides.
    Args:
        concat_slides (bool): Specifies whether to concatenate all slide text into one document.
            - If True, slide texts will be joined together as a single string.
            - If False, each slide's text will be stored as a separate entry in a list.
            Set to True by default.
        slide_separator (str): Separator used to join slides' text content.
            Only used when `concat_slides=True`. Default is "\n".
        Refer to https://python-pptx.readthedocs.io/en/latest/ for more information.
    """

    def __init__(
        self,
        *args: Any,
        concat_slides: bool = True,
        slide_separator: str = "\n",
        **kwargs: Any
    ) -> None:
        """Init params."""
        super().__init__(*args, **kwargs)
        self._concat_slides = concat_slides
        self._slide_separator = slide_separator

    def _init_parser(self) -> Dict:
        """Init parser."""
        return {}

    def parse_file(self, file: Path, errors: str = "ignore") -> Union[str, List[str]]:
        r"""
        Parse a .pptx file and extract text from each slide.
        Args:
            file (Path): Path to the .pptx file.
            errors (str): Error handling policy ('ignore' by default).
        Returns:
            Union[str, List[str]]: Concatenated text if concat_slides is True,
            otherwise a list of slide texts.
        """

        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("pptx module is required to read .PPTX files.")

        try:
            presentation = Presentation(file)
            slide_texts=[]

            # Iterate over each slide in the presentation
            for slide in presentation.slides:
                blocks: List[str] = []

                # Iterate over each shape in the slide
                for shape in slide.shapes:
                    blocks.extend(self._shape_text(shape))

                slide_texts.append("\n".join(blocks).strip())

            if self._concat_slides:
                return self._slide_separator.join(slide_texts)
            else:
                return slide_texts

        except Exception as e:
            raise e

    def _shape_text(self, shape: Any) -> List[str]:
        """Return the blocks of text a shape carries, in reading order.

        A group carries no text of its own and neither does the graphic frame a
        table lives on, so `shape.text` alone misses both entirely. A group is
        the only shape with its own `shapes` collection, which is what
        identifies it here without importing the enum.
        """
        if hasattr(shape, "shapes"):
            grouped: List[str] = []
            for grouped_shape in shape.shapes:
                grouped.extend(self._shape_text(grouped_shape))
            return grouped

        if getattr(shape, "has_table", False):
            return [self._table_to_markdown(shape.table)]

        text = getattr(shape, "text", "") or ""
        text = text.strip()
        return [text] if text else []

    @staticmethod
    def _escape_cell(text: str) -> str:
        """Make a cell's text safe between the pipes of a GFM row.

        A cell with more than one paragraph carries a line break, which would end
        the row in the middle of it, and a pipe would add a column. The backslash
        run in front of a pipe is doubled first, so a backslash the cell already
        contains cannot consume the escape.
        """
        text = _CELL_BREAK.sub(" ", text.strip())
        return _CELL_PIPE.sub(lambda match: match.group(1) * 2 + r"\|", text)

    @staticmethod
    def _table_to_markdown(table: Any) -> str:
        """Render a slide table as a GFM table, the form `tableize` also emits."""
        rows = [[PPTXParser._escape_cell(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            return ""
        lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * len(rows[0])]
        lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        return "\n".join(lines)