"""Turn a document into numbered lines a model can cite and code can check.

``extract_document_lines`` reads DOCX, XLSX and PDF with format-specific
readers so the structure survives: a DOCX keeps its paragraph and table
order, a table row becomes one line with its cells joined by ``" | "``, an
XLSX row carries its sheet and first cell (``(Sheet1!A5)``) and a PDF line
its page (``(p.3)``). Plain text and CSV are read directly; every other
format goes through the configured parser (``DOC_PARSER_ENGINE``), whose
Markdown tables are turned into rows the same way. Blank lines are dropped,
and nothing is ever cut: the caller decides whether the text fits.

The same untrusted-input guards as ``parse_document_bytes`` apply: an
extension whitelist, the parse byte cap and the zip-bomb check.
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from docsgpt.parser.file.constants import SUPPORTED_SOURCE_EXTENSIONS
from docsgpt.utils import safe_filename

logger = logging.getLogger(__name__)

# Bumped whenever the line format changes, so cached line text from an
# older format is never served.
FORMAT_VERSION = 1

CELL_SEPARATOR = " | "

_TEXT_SUFFIXES = frozenset({".txt", ".md", ".mdx", ".rst", ".log", ".json", ".xml", ".yaml", ".yml"})
_CSV_SUFFIXES = frozenset({".csv", ".tsv"})
_XLSX_SUFFIXES = frozenset({".xlsx", ".xlsm"})

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

_MD_ROW = re.compile(r"^\s*\|.*\|\s*$")
_MD_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_PLAIN_SHEET_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")


class _Lines:
    """Accumulates lines and the table spans they form."""

    def __init__(self) -> None:
        self.lines: List[str] = []
        self.tables: List[Dict[str, Any]] = []
        self._table: Optional[Dict[str, Any]] = None

    def text(self, value: str, marker: Optional[str] = None) -> None:
        """Add a block of text, one line per non-blank line in it."""
        self.end_table()
        for line in value.splitlines():
            self._add(line, marker)

    def row(self, cells: List[str], marker: Optional[str] = None) -> None:
        """Add one table row; the first row of a table is its header."""
        cells = [" ".join(cell.split()) for cell in cells]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            return
        if not self._add(CELL_SEPARATOR.join(cells), marker, strip=False):
            return
        line_no = len(self.lines)
        if self._table is None:
            self._table = {"first_line": line_no, "last_line": line_no, "header": cells}
            self.tables.append(self._table)
        self._table["last_line"] = line_no

    def end_table(self) -> None:
        """Close the open table, so the next row starts a new one."""
        self._table = None

    def _add(self, line: str, marker: Optional[str], strip: bool = True) -> bool:
        line = line.strip() if strip else line.rstrip()
        if not line.strip():
            return False
        self.lines.append(f"({marker}) {line}" if marker else line)
        return True

    def result(self) -> Dict[str, Any]:
        return {"lines": self.lines, "tables": self.tables}


def number_lines(document_number: int, lines: Iterable[str]) -> str:
    """Prefix each line with its ``[D#:L#]`` reference and join them.

    Args:
        document_number: The document's 1-based number in the run (``D#``).
        lines: The document's lines, as ``extract_document_lines`` returns them.

    Returns:
        The numbered text, one line per input line.
    """
    return "\n".join(f"[D{document_number}:L{index}] {line}" for index, line in enumerate(lines, start=1))


def _max_input_bytes() -> int:
    from docsgpt.parser.document_reader import _max_input_bytes as reader_cap

    return reader_cap()


def extract_document_lines(data: bytes, filename: str, mime_type: Optional[str] = None) -> Dict[str, Any]:
    """Read a document into lines; never truncates.

    Args:
        data: The document's bytes (untrusted).
        filename: Its display name; the extension picks the reader.
        mime_type: Its stored mime type; a ``text/*`` document with no
            known extension is read as plain text.

    Returns:
        ``{"lines": [...], "tables": [{"first_line", "last_line", "header"}]}``
        with 1-based line numbers, or ``{"error": message}``.
    """
    from docsgpt.parser.document_reader import _reject_zip_bomb

    safe_name = safe_filename(filename) or "document"
    suffix = os.path.splitext(safe_name)[1].lower()
    is_text = suffix in _TEXT_SUFFIXES or (
        suffix not in SUPPORTED_SOURCE_EXTENSIONS and (mime_type or "").startswith("text/")
    )
    if not is_text and suffix not in SUPPORTED_SOURCE_EXTENSIONS:
        return {"error": f"unsupported file type '{suffix or filename}'."}

    cap = _max_input_bytes()
    if len(data) > cap:
        return {"error": f"input document is too large: {len(data)} bytes exceeds the {cap}-byte cap."}
    bomb = _reject_zip_bomb(data, suffix)
    if bomb is not None:
        return {"error": bomb}

    try:
        if is_text:
            return _text_lines(data.decode("utf-8", errors="replace"))
        if suffix in _CSV_SUFFIXES:
            return _csv_lines(data, delimiter="\t" if suffix == ".tsv" else ",")
        if suffix == ".docx":
            return _docx_lines(data)
        if suffix in _XLSX_SUFFIXES:
            return _xlsx_lines(data)
        if suffix == ".pdf":
            return _pdf_lines(data, safe_name)
        return _text_lines(_parse_with_configured_parser(data, safe_name, suffix))
    except Exception as exc:
        logger.warning("extract_document_lines: could not read %s: %s", suffix, type(exc).__name__, exc_info=True)
        return {"error": f"could not read the {suffix.lstrip('.') or 'file'} document."}


def _text_lines(text: str, marker: Optional[str] = None, out: Optional[_Lines] = None) -> Dict[str, Any]:
    """Plain or Markdown text: Markdown table rows become cell rows, separators are dropped."""
    out = out or _Lines()
    for line in text.splitlines():
        if _MD_ROW.match(line):
            if _MD_SEPARATOR.match(line):
                continue
            out.row(_markdown_cells(line), marker)
        else:
            out.text(line, marker)
    out.end_table()
    return out.result()


def _markdown_cells(line: str) -> List[str]:
    body = line.strip()[1:-1].replace("\\|", "\x00")
    return [cell.replace("\x00", "|").strip() for cell in body.split("|")]


def _csv_lines(data: bytes, delimiter: str) -> Dict[str, Any]:
    out = _Lines()
    text = data.decode("utf-8-sig", errors="replace")
    for cells in csv.reader(io.StringIO(text), delimiter=delimiter):
        out.row(cells)
    return out.result()


def _docx_lines(data: bytes) -> Dict[str, Any]:
    """Walk the DOCX body in order: each paragraph is a line, each table row is a line."""
    from defusedxml import ElementTree

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    body = root.find(f"{_W}body")
    out = _Lines()
    if body is not None:
        _docx_walk(body, out)
    return out.result()


def _docx_walk(element: Any, out: _Lines) -> None:
    for child in element:
        if child.tag == f"{_W}p":
            out.text(_docx_paragraph_text(child))
        elif child.tag == f"{_W}tbl":
            out.end_table()
            for row in child.findall(f"{_W}tr"):
                out.row([_docx_cell_text(cell) for cell in row.findall(f"{_W}tc")])
            out.end_table()
        elif child.tag == f"{_W}sdt":
            content = child.find(f"{_W}sdtContent")
            if content is not None:
                _docx_walk(content, out)


def _docx_paragraph_text(paragraph: Any) -> str:
    parts: List[str] = []
    for node in paragraph.iter():
        if node.tag == f"{_W}t" and node.text:
            parts.append(node.text)
        elif node.tag == f"{_W}tab":
            parts.append("\t")
        elif node.tag in (f"{_W}br", f"{_W}cr"):
            parts.append("\n")
        elif node.tag == f"{_W}noBreakHyphen":
            parts.append("-")
    return "".join(parts)


def _docx_cell_text(cell: Any) -> str:
    return " ".join(_docx_paragraph_text(p) for p in cell.iter(f"{_W}p")).strip()


def _xlsx_lines(data: bytes) -> Dict[str, Any]:
    """Each sheet is one table; a row's marker is its sheet and the sheet's first used column."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = _Lines()
    try:
        for sheet in workbook.worksheets:
            # The declared <dimension> comes from the file; trusting it pads
            # every row to its width (up to 16,384 columns).
            sheet.reset_dimensions()
            rows = []
            first_col: Optional[int] = None
            for row_number, values in enumerate(sheet.iter_rows(min_row=1, min_col=1, values_only=True), start=1):
                cells = [_cell_text(value) for value in values]
                used = [index for index, cell in enumerate(cells) if cell]
                if not used:
                    continue
                first_col = used[0] if first_col is None else min(first_col, used[0])
                rows.append((row_number, cells))
            name = sheet.title if _PLAIN_SHEET_NAME.match(sheet.title) else "'" + sheet.title.replace("'", "''") + "'"
            for row_number, cells in rows:
                out.row(cells[first_col:], f"{name}!{get_column_letter(first_col + 1)}{row_number}")
            out.end_table()
    finally:
        workbook.close()
    return out.result()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def _pdf_lines(data: bytes, safe_name: str) -> Dict[str, Any]:
    """Text-layer lines with a page marker; a PDF with no text layer goes to the configured (OCR) parser."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt("")
    out = _Lines()
    for page_number, page in enumerate(reader.pages, start=1):
        out.text(page.extract_text() or "", f"p.{page_number}")
    if out.lines:
        return out.result()
    text = _parse_with_configured_parser(data, safe_name, ".pdf")
    if "\f" not in text:
        return _text_lines(text)
    out = _Lines()
    for page_number, page_text in enumerate(text.split("\f"), start=1):
        _text_lines(page_text, f"p.{page_number}", out)
    return out.result()


def _parse_with_configured_parser(data: bytes, safe_name: str, suffix: str) -> str:
    """Text from the parser ``DOC_PARSER_ENGINE`` picks for ``suffix`` (OCR as configured)."""
    from docsgpt.parser.document_reader import (
        _effective_engine,
        _parse_to_text,
        _pick_parser,
        _resolve_ocr_enabled,
    )

    parser = _pick_parser(suffix, ocr_enabled=_resolve_ocr_enabled("auto"), engine=_effective_engine("auto"))
    tmp_dir = tempfile.mkdtemp(prefix="doclines-")
    tmp_path = Path(tmp_dir) / safe_name
    try:
        tmp_path.write_bytes(data)
        if parser is None:
            return tmp_path.read_text(errors="ignore")
        return _parse_to_text(parser, tmp_path)
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
            os.rmdir(tmp_dir)
        except OSError:
            logger.warning("extract_document_lines: temp cleanup failed for %s", tmp_dir, exc_info=True)
