"""Smoke test for ``docsgpt.worker.attachment_worker``.

The happy path parses an uploaded file and inserts a row into
``attachments``. We mock the parser boundary (``StorageCreator.get_storage``
returns a storage whose ``process_file`` produces a pre-built Document)
but let the PG insert run against the ephemeral ``pg_conn`` so we can
assert one concrete row is visible after the task returns.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest

from docsgpt.parser.schema.base import Document
from docsgpt.storage.db.repositories.attachments import AttachmentsRepository


@pytest.mark.unit
class TestAttachmentWorker:
    def test_inserts_row_in_attachments(
        self, pg_conn, patch_worker_db, task_self, monkeypatch
    ):
        from docsgpt import worker

        fake_doc = Document(
            text="hello world",
            extra_info={"transcript_language": "en"},
        )

        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.return_value = fake_doc
        monkeypatch.setattr(
            worker.StorageCreator, "get_storage", lambda: fake_storage
        )

        # Stub the parser selection so the docling import path isn't taken.
        monkeypatch.setattr(
            worker, "get_default_file_extractor", lambda ocr_enabled=False, pdf_text_fast_path=False: {}
        )

        file_info = {
            "filename": "notes.txt",
            "attachment_id": "507f1f77bcf86cd799439011",
            "path": "uploads/user1/notes.txt",
            "metadata": {"source": "chat"},
        }

        result = worker.attachment_worker(task_self, file_info, "user1")

        assert result["filename"] == "notes.txt"
        assert result["token_count"] > 0
        # Parser metadata (``transcript_*``) should have been merged in.
        assert result["metadata"]["transcript_language"] == "en"
        assert result["metadata"]["source"] == "chat"

        # Row should be resolvable by the caller-visible handle stored in
        # ``legacy_mongo_id``.
        row = AttachmentsRepository(pg_conn).get_by_legacy_id(
            file_info["attachment_id"], "user1"
        )
        assert row is not None, "attachment_worker should insert a row"
        assert row["filename"] == "notes.txt"
        assert row["upload_path"] == "uploads/user1/notes.txt"
        assert row["content"] == "hello world"
        assert row["user_id"] == "user1"

    def test_parse_failure_stores_failure_row_and_tells_the_user(
        self, pg_conn, patch_worker_db, task_self, monkeypatch
    ):
        """A parse failure must fail loudly, not store the error as content.

        Regression (prod 2026-08-05): docling's PDF layout stage died, the
        parser returned its own traceback as the "document", and the worker
        stored it and published ``attachment.completed`` — so the upload
        looked fine and the model was handed an error message as the PDF.
        The failure now leaves a queryable row whose ``metadata.extraction``
        records the error; ``content`` stays NULL so the model can never
        read a traceback as the document.
        """
        from docsgpt import worker
        from docsgpt.parser.file.base_parser import DocumentParseError

        published: list[tuple[str, dict]] = []
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, event, payload, **kw: published.append((event, payload)),
        )

        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = DocumentParseError(
            "Failed to parse scan.pdf with docling: Conversion failed for: "
            "scan.pdf with status: failure. Errors: InvalidCxxCompiler"
        )
        monkeypatch.setattr(
            worker.StorageCreator, "get_storage", lambda: fake_storage
        )
        monkeypatch.setattr(
            worker, "get_default_file_extractor", lambda ocr_enabled=False, pdf_text_fast_path=False: {}
        )

        file_info = {
            "filename": "scan.pdf",
            "attachment_id": "507f1f77bcf86cd799439012",
            "path": "uploads/user1/scan.pdf",
            "metadata": {"source": "chat"},
        }

        with pytest.raises(DocumentParseError):
            worker.attachment_worker(task_self, file_info, "user1")

        # The error lives in metadata, never in content — an attachment whose
        # text is a traceback is worse than no attachment, because the model
        # will read it.
        row = AttachmentsRepository(pg_conn).get_by_legacy_id(
            file_info["attachment_id"], "user1"
        )
        assert row is not None, "a failed parse must leave a queryable row"
        assert row["content"] is None
        extraction = row["metadata"]["extraction"]
        assert extraction["status"] == "failed"
        assert "InvalidCxxCompiler" in extraction["error"]

        # The user is told it failed, and never told it completed.
        events = [event for event, _ in published]
        assert "attachment.failed" in events
        assert "attachment.completed" not in events

    @pytest.mark.parametrize("task_name", ["store_attachment", "ingest"])
    def test_parse_failure_is_not_retried(self, task_name):
        """A parse failure is deterministic; retrying only multiplies noise.

        Both parsing entry points matter: now that the parser raises instead of
        returning its traceback as content, an unguarded task would turn one
        unreadable upload into a retry loop of identical failures.
        """
        from docsgpt.api.user import tasks as user_tasks
        from docsgpt.parser.file.base_parser import DocumentParseError

        task = getattr(user_tasks, task_name)
        assert DocumentParseError in task.dont_autoretry_for


@pytest.mark.unit
class TestBoundedAttachmentCopy:
    """``_bounded_attachment_copy`` head-truncates oversized line-oriented
    text attachments to a temp copy before parsing.

    The parsed content is capped at ~250k chars downstream anyway, so bytes
    beyond the cap only cost parse time and memory. Local storage hands the
    canonical stored file to the processor, so truncation must never happen
    in place.
    """

    def _write(self, tmp_path, name: str, data: bytes):
        path = tmp_path / name
        path.write_bytes(data)
        return path

    def test_oversized_csv_is_copied_and_truncated(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 1024)
        original = self._write(
            tmp_path, "big.csv", b"".join(b"%d,%d\n" % (i, i) for i in range(1000))
        )
        original_size = original.stat().st_size

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert is_temp is True
        assert parse_path != str(original)
        assert parse_path.endswith(".csv")
        copied = Path(parse_path).read_bytes()
        assert 0 < len(copied) <= 1024
        assert copied.endswith(b"\n"), "must cut on a line boundary"
        # The stored original must be untouched.
        assert original.stat().st_size == original_size
        Path(parse_path).unlink()

    def test_small_file_returned_as_is(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 1024)
        original = self._write(tmp_path, "small.csv", b"a,b\n1,2\n")

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert parse_path == str(original)
        assert is_temp is False

    def test_non_text_suffix_is_never_truncated(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 64)
        original = self._write(tmp_path, "doc.pdf", b"%PDF-1.7 " + b"x" * 500)

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert parse_path == str(original)
        assert is_temp is False

    def test_cap_zero_disables_truncation(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 0)
        original = self._write(tmp_path, "big.csv", b"1,2\n" * 1000)

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert parse_path == str(original)
        assert is_temp is False

    def test_single_line_without_newline_falls_back_to_hard_cut(
        self, tmp_path, monkeypatch
    ):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 256)
        original = self._write(tmp_path, "oneline.txt", b"x" * 5000)

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert is_temp is True
        data = Path(parse_path).read_bytes()
        assert len(data) == 256
        Path(parse_path).unlink()

    def test_leading_newline_does_not_collapse_the_copy(self, tmp_path, monkeypatch):
        """A window whose only newline sits at byte 0 must keep its content.

        ``rfind`` returns 0 for this shape, so cutting at that boundary would
        write a one-byte copy and throw the attachment away — the partial
        final line is the better trade.
        """
        from docsgpt import worker
        from docsgpt.core.settings import settings

        monkeypatch.setattr(settings, "ATTACHMENT_TEXT_MAX_BYTES", 256)
        original = self._write(tmp_path, "leading.log", b"\n" + b"x" * 5000)

        parse_path, is_temp = worker._bounded_attachment_copy(str(original))

        assert is_temp is True
        data = Path(parse_path).read_bytes()
        assert len(data) == 256
        Path(parse_path).unlink()


@pytest.mark.unit
class TestAttachmentZipBombGuard:
    """``_reject_attachment_zip_bomb`` brings the ingest-path zip-bomb guard to
    the attachment path (which previously had none): a zip-container attachment
    that declares too many entries / too much inner data is rejected before any
    parser touches it, with a non-retryable error."""

    def _make_xlsx(self, path: Path, rows: int = 20):
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        for i in range(rows):
            ws.append([i, i * 2, i * 3])
        wb.save(str(path))

    def test_rejects_when_inner_size_exceeds_cap(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        path = tmp_path / "book.xlsx"
        self._make_xlsx(path)
        monkeypatch.setattr(settings, "DOCUMENT_MAX_DECOMPRESSED_BYTES", 100)

        with pytest.raises(worker.AttachmentRejectedError):
            worker._reject_attachment_zip_bomb(str(path))

    def test_rejects_when_too_many_entries(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        path = tmp_path / "book.xlsx"
        self._make_xlsx(path)
        monkeypatch.setattr(settings, "DOCUMENT_MAX_ARCHIVE_ENTRIES", 1)

        with pytest.raises(worker.AttachmentRejectedError):
            worker._reject_attachment_zip_bomb(str(path))

    def test_allows_reasonable_archive(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        path = tmp_path / "book.xlsx"
        self._make_xlsx(path)
        monkeypatch.setattr(settings, "DOCUMENT_MAX_DECOMPRESSED_BYTES", 300 * 1024 * 1024)
        monkeypatch.setattr(settings, "DOCUMENT_MAX_ARCHIVE_ENTRIES", 10000)

        # Must not raise.
        worker._reject_attachment_zip_bomb(str(path))

    def test_non_container_suffix_is_ignored(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        path = tmp_path / "notes.txt"
        path.write_bytes(b"x" * 5000)
        monkeypatch.setattr(settings, "DOCUMENT_MAX_DECOMPRESSED_BYTES", 1)

        # Text files are not zip containers — never inspected, never rejected.
        worker._reject_attachment_zip_bomb(str(path))

    def test_corrupt_zip_is_left_to_the_parser(self, tmp_path, monkeypatch):
        from docsgpt import worker
        from docsgpt.core.settings import settings

        path = tmp_path / "broken.xlsx"
        path.write_bytes(b"not a real zip")
        monkeypatch.setattr(settings, "DOCUMENT_MAX_DECOMPRESSED_BYTES", 1)

        # BadZipFile → return quietly; the format parser surfaces a clean error.
        worker._reject_attachment_zip_bomb(str(path))


@pytest.mark.unit
class TestAttachmentTypeGuard:
    """A binary with no parser must fail rather than be read as text.

    ``SimpleDirectoryReader`` falls through to a plain-text ``open()`` for a
    suffix it has no parser for, which is how a phone-uploaded ``.mp4`` was
    once stored as 100k tokens of binary with ``extraction.status == "ok"``.
    The same fallthrough is what makes a ``.py`` or ``.log`` attachment work,
    so the guard admits those on content.
    """

    def test_binary_without_a_parser_fails_instead_of_being_read_as_text(
        self, pg_conn, patch_worker_db, task_self, monkeypatch, tmp_path
    ):
        from docsgpt import worker

        local_path = tmp_path / "clip.mp4"
        local_path.write_bytes(b"\x00\x00\x00\x18ftypisom\x00\x00\x02\x00isomavc1")

        events = []
        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = lambda path, callback: callback(
            str(local_path)
        )
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        monkeypatch.setattr(
            worker,
            "get_default_file_extractor",
            lambda ocr_enabled=False, pdf_text_fast_path=False: {},
        )
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, name, payload, **kwargs: events.append((name, payload)),
        )

        file_info = {
            "filename": "clip.mp4",
            "attachment_id": "507f1f77bcf86cd799439012",
            "path": "uploads/user1/attachments/clip.mp4",
            "metadata": {"source": "chat"},
        }

        with pytest.raises(
            worker.AttachmentRejectedError, match=r"Unsupported file type: \.mp4"
        ):
            worker.attachment_worker(task_self, file_info, "user1")

        failed = [payload for name, payload in events if name == "attachment.failed"]
        assert failed and failed[0]["error"] == "Unsupported file type: .mp4"

    def test_suffix_the_loaded_extractor_cannot_parse_is_refused(
        self, pg_conn, patch_worker_db, task_self, monkeypatch, tmp_path
    ):
        """The guard judges against the parser table actually loaded.

        The route admits ``.vtt`` by name, but a docling-less install has no
        VTT parser, so a binary file named ``.vtt`` would otherwise be opened
        as plain text here.
        """
        from docsgpt import worker

        local_path = tmp_path / "subs.vtt"
        local_path.write_bytes(b"\x00\x01\x02\x03" + bytes(range(256)))

        events = []
        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = lambda path, callback: callback(
            str(local_path)
        )
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        # A trimmed parser table: .png is handled, .vtt is not.
        monkeypatch.setattr(
            worker,
            "get_default_file_extractor",
            lambda ocr_enabled=False, pdf_text_fast_path=False: {".png": object()},
        )
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, name, payload, **kwargs: events.append((name, payload)),
        )

        file_info = {
            "filename": "subs.vtt",
            "attachment_id": "507f1f77bcf86cd799439014",
            "path": "uploads/user1/attachments/subs.vtt",
            "metadata": {"source": "chat"},
        }

        with pytest.raises(
            worker.AttachmentRejectedError, match=r"Unsupported file type: \.vtt"
        ):
            worker.attachment_worker(task_self, file_info, "user1")

        failed = [payload for name, payload in events if name == "attachment.failed"]
        assert failed and failed[0]["error"] == "Unsupported file type: .vtt"

    def test_scanned_pdf_completes_without_text_for_models_that_read_it_natively(
        self, pg_conn, patch_worker_db, task_self, monkeypatch
    ):
        """A PDF with no text layer is still a usable attachment.

        Models that read PDFs natively (or as page images) are sent the
        stored file, not its extracted text, so failing the upload only kept
        the file from a model that could read it (prod 2026-09-12: a paying
        user's scanned club bylaws). The row is kept with ``status: no_text``
        — text inlining still skips it — and the user is told it completed.
        """
        from docsgpt import worker
        from docsgpt.parser.file.base_parser import NoTextLayerError

        published: list[tuple[str, dict]] = []
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, event, payload, **kw: published.append((event, payload)),
        )
        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = NoTextLayerError(
            "bylaws.pdf appears to be a scanned PDF (no text layer), and "
            "PDFParser extracted almost nothing"
        )
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        monkeypatch.setattr(
            worker, "get_default_file_extractor", lambda ocr_enabled=False, pdf_text_fast_path=False: {}
        )

        file_info = {
            "filename": "bylaws.pdf",
            "attachment_id": "507f1f77bcf86cd799439021",
            "path": "uploads/user1/attachments/bylaws.pdf",
            "metadata": {"source": "chat"},
        }

        result = worker.attachment_worker(task_self, file_info, "user1")

        assert result["token_count"] == 0
        assert result["mime_type"] == "application/pdf"
        extraction = result["metadata"]["extraction"]
        assert extraction["status"] == "no_text"
        assert "scanned PDF" in extraction["reason"]

        row = AttachmentsRepository(pg_conn).get_by_legacy_id(
            file_info["attachment_id"], "user1"
        )
        assert row["upload_path"] == file_info["path"]
        assert row["content"] == ""
        assert row["metadata"]["extraction"]["status"] == "no_text"

        events = [event for event, _ in published]
        assert "attachment.failed" not in events
        completed = [payload for event, payload in published if event == "attachment.completed"]
        assert completed[0]["extraction_status"] == "no_text"
        assert completed[0]["mime_type"] == "application/pdf"

    def test_no_text_on_a_type_models_cannot_read_natively_still_fails(
        self, pg_conn, patch_worker_db, task_self, monkeypatch
    ):
        """Only PDFs and images have a native path; anything else without text stays a failure."""
        from docsgpt import worker
        from docsgpt.parser.file.base_parser import NoTextLayerError

        published: list[tuple[str, dict]] = []
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, event, payload, **kw: published.append((event, payload)),
        )
        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = NoTextLayerError("notes.docx has no text")
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        monkeypatch.setattr(
            worker, "get_default_file_extractor", lambda ocr_enabled=False, pdf_text_fast_path=False: {}
        )

        file_info = {
            "filename": "notes.docx",
            "attachment_id": "507f1f77bcf86cd799439022",
            "path": "uploads/user1/attachments/notes.docx",
            "metadata": {"source": "chat"},
        }

        with pytest.raises(NoTextLayerError):
            worker.attachment_worker(task_self, file_info, "user1")

        events = [event for event, _ in published]
        assert "attachment.failed" in events
        assert "attachment.completed" not in events

    def test_tiff_is_stored_as_png_so_providers_accept_it(
        self, pg_conn, patch_worker_db, task_self, monkeypatch
    ):
        """OpenAI and Anthropic reject TIFF, so the row must point at a PNG copy."""
        import io

        from PIL import Image

        from docsgpt import worker

        buf = io.BytesIO()
        Image.new("RGB", (32, 20), color=(10, 120, 200)).save(buf, format="TIFF")
        tiff_bytes = buf.getvalue()

        published: list[tuple[str, dict]] = []
        saved: dict[str, bytes] = {}

        def _save(data, path, **kwargs):
            saved[path] = data.read()
            return {"storage_type": "local"}

        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.return_value = Document(text="", extra_info={})
        fake_storage.get_file.side_effect = lambda path: io.BytesIO(tiff_bytes)
        fake_storage.save_file.side_effect = _save
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        monkeypatch.setattr(
            worker, "get_default_file_extractor", lambda ocr_enabled=False, pdf_text_fast_path=False: {}
        )
        monkeypatch.setattr(
            worker,
            "publish_user_event",
            lambda user, event, payload, **kw: published.append((event, payload)),
        )

        file_info = {
            "filename": "fax.tiff",
            "attachment_id": "507f1f77bcf86cd799439023",
            "path": "uploads/user1/attachments/abc/fax.tiff",
            "metadata": {"source": "chat"},
        }

        result = worker.attachment_worker(task_self, file_info, "user1")

        png_path = "uploads/user1/attachments/abc/fax.png"
        assert list(saved) == [png_path]
        assert Image.open(io.BytesIO(saved[png_path])).format == "PNG"
        assert result["path"] == png_path
        assert result["mime_type"] == "image/png"
        assert result["metadata"]["image_conversion"] == {
            "from": "image/tiff",
            "to": "image/png",
            "frames": 1,
        }

        row = AttachmentsRepository(pg_conn).get_by_legacy_id(
            file_info["attachment_id"], "user1"
        )
        assert row["filename"] == "fax.tiff"
        assert row["upload_path"] == png_path
        assert row["mime_type"] == "image/png"

        completed = [payload for event, payload in published if event == "attachment.completed"]
        assert completed[0]["mime_type"] == "image/png"

    def test_text_without_a_parser_is_parsed(
        self, pg_conn, patch_worker_db, task_self, monkeypatch, tmp_path
    ):
        from docsgpt import worker

        local_path = tmp_path / "server.log"
        local_path.write_text("2026-09-02 ERROR boom\n", encoding="utf-8")

        fake_storage = MagicMock(name="storage")
        fake_storage.process_file.side_effect = lambda path, callback: callback(
            str(local_path)
        )
        monkeypatch.setattr(worker.StorageCreator, "get_storage", lambda: fake_storage)
        monkeypatch.setattr(
            worker,
            "get_default_file_extractor",
            lambda ocr_enabled=False, pdf_text_fast_path=False: {},
        )

        file_info = {
            "filename": "server.log",
            "attachment_id": "507f1f77bcf86cd799439013",
            "path": "uploads/user1/attachments/server.log",
            "metadata": {"source": "chat"},
        }

        result = worker.attachment_worker(task_self, file_info, "user1")

        assert result["filename"] == "server.log"
        assert result["token_count"] > 0
