"""Generate the synthetic many-attachments fixtures described in scenarios.yaml.

Nothing large is committed: tests (and end-to-end scripts) call ``generate`` to
write the files into a temporary directory. Output is deterministic for a given
scenario, profile and library versions (seeded RNGs, invariant PDFs, fixed zip
timestamps), so duplicates hash-equal and reruns are stable.

Usage from a test::

    from tests.fixtures.many_attachments import generate as gen

    manifest = gen.generate(tmp_path, ["RC-01"])          # small profile
    rows = gen.attachment_records(gen.get_scenario("RC-04"))  # no files, full-size targets
    for caps in gen.capability_matrix(gen.get_scenario("RC-04")):
        for conv in gen.conversations(gen.get_scenario("RC-04")):
            for turn in conv:
                exp = gen.expectations_for(gen.get_scenario("RC-04"), turn, caps)

From a shell::

    python -m tests.fixtures.many_attachments.generate --out /tmp/ma RC-01 V1-01 [--full]

Profiles: ``small`` (default) caps text files at ~3k tokens, spreadsheets at
300 rows, scans at 2 pages and shrinks images, so a full run takes seconds.
``full`` writes the production-shaped sizes (100k-400k-token PDFs, 200k-row
CSVs, 400-page scans) and takes minutes.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import os
import random
import re
import shutil
import textwrap
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import yaml

from tests.fixtures.many_attachments import corpus

HERE = Path(__file__).resolve().parent
SCENARIOS_PATH = HERE / "scenarios.yaml"
BASE_SEED = 20261001
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
FIXED_DT = datetime(2026, 1, 1, tzinfo=timezone.utc)

PROFILES: Dict[str, Dict[str, Any]] = {
    "small": {"token_cap": 3000, "rows_cap": 300, "image_scale": 0.4, "scan_pages_cap": 2, "pages_cap": 6},
    "full": {"token_cap": None, "rows_cap": None, "image_scale": 1.0, "scan_pages_cap": None, "pages_cap": None},
}

TEXT_KINDS = {
    "invoice_pdf", "long_pdf", "exam_pdf", "report_pdf", "paper_pdf", "statement_pdf", "docx_writeup",
    "notebook_html", "log_txt", "python_script", "xml_dump", "epub_book", "sl_ordinance_pdf", "sl_contract",
    "sl_certificate_pdf", "sl_filing_pdf", "sl_extract_text", "nextjs_zip", "markdown_zip",
}
SCAN_KINDS = {"scan_pdf", "sl_scan_pdf"}
IMAGE_KINDS = {"screenshot_image", "photo_image", "quiz_image", "map_image"}
SHEET_KINDS = {"xlsx_workbook", "portfolio_xlsx", "dues_xlsx", "csv_table"}

MIME = {
    ".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".webp": "image/webp", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".csv": "text/csv", ".html": "text/html", ".txt": "text/plain", ".py": "text/x-python",
    ".xml": "application/xml", ".epub": "application/epub+zip", ".zip": "application/zip",
    ".heic": "image/heic", ".md": "text/markdown",
}


def mime_for(name: str) -> str:
    """Return the MIME type for a file name (case-insensitive extension)."""
    return MIME.get(Path(name).suffix.lower(), "application/octet-stream")


# ---------------------------------------------------------------------------
# Scenario model
# ---------------------------------------------------------------------------


@dataclass
class FileSpec:
    """One synthetic file of a scenario.

    Attributes:
        scenario: Scenario id the file belongs to.
        group: File group name from scenarios.yaml.
        index: 0-based index within the group.
        name: Final file name (extension included).
        kind: Writer kind.
        tokens: Full-profile token target (None for images/spreadsheets).
        rows: Full-profile row target for spreadsheets/CSV.
        params: Remaining group parameters.
        duplicate_of: Key of the source file for ``kind: duplicate``.
    """

    scenario: str
    group: str
    index: int
    name: str
    kind: str
    tokens: Optional[int] = None
    rows: Optional[int] = None
    params: Dict[str, Any] = field(default_factory=dict)
    duplicate_of: Optional[str] = None

    @property
    def key(self) -> str:
        """Stable key, ``group[index]``."""
        return f"{self.group}[{self.index}]"

    @property
    def rel_path(self) -> str:
        """Path relative to the scenario output directory."""
        return f"{self.group}/{self.name}"

    def seed(self, salt: str = "") -> int:
        """Deterministic seed for this file (independent of profile)."""
        src = self.params.get("_seed_scope", self.scenario)
        h = hashlib.sha256(f"{BASE_SEED}:{src}:{self.group}:{self.index}:{salt}".encode()).hexdigest()
        return int(h[:16], 16)


@dataclass
class Turn:
    """One user turn of a scenario conversation."""

    conversation: int
    index: int
    attach: List[str]
    uploading: int = 0
    prompt: str = ""
    extra: Dict[str, Any] = field(default_factory=dict)


_SCENARIOS: Optional[Dict[str, Any]] = None


def load_scenarios(path: Path = SCENARIOS_PATH) -> Dict[str, Any]:
    """Load and cache scenarios.yaml."""
    global _SCENARIOS
    if path != SCENARIOS_PATH:
        return yaml.safe_load(path.read_text(encoding="utf-8"))
    if _SCENARIOS is None:
        _SCENARIOS = yaml.safe_load(path.read_text(encoding="utf-8"))
    return _SCENARIOS


def scenario_ids() -> List[str]:
    """Return every scenario id in file order."""
    return [s["id"] for s in load_scenarios()["scenarios"]]


def get_scenario(scenario_id: str) -> Dict[str, Any]:
    """Return one scenario dict by id."""
    for s in load_scenarios()["scenarios"]:
        if s["id"] == scenario_id:
            return s
    raise KeyError(scenario_id)


def _per_file(value: Any, count: int, rng: random.Random) -> List[Optional[int]]:
    """Expand an int / [min, max] range / explicit list into one value per file."""
    if value is None:
        return [None] * count
    if isinstance(value, int):
        return [value] * count
    if isinstance(value, list) and len(value) == count:
        return [int(v) for v in value]
    if isinstance(value, list) and len(value) == 2:
        return [rng.randint(int(value[0]), int(value[1])) for _ in range(count)]
    raise ValueError(f"cannot expand {value!r} to {count} files")


def _file_groups(scenario: Dict[str, Any]) -> Tuple[str, List[Dict[str, Any]]]:
    src = scenario.get("files_from")
    if src:
        return src, get_scenario(src)["files"]
    return scenario["id"], scenario["files"]


def file_specs(scenario: Dict[str, Any]) -> List[FileSpec]:
    """Expand a scenario's file groups into ``FileSpec`` objects (sources before duplicates)."""
    seed_scope, groups = _file_groups(scenario)
    specs: List[FileSpec] = []
    by_key: Dict[str, FileSpec] = {}
    for g in groups:
        kind = g["kind"]
        if kind == "duplicate":
            continue
        count = int(g.get("count", 1))
        rng = random.Random(int(hashlib.sha256(f"{BASE_SEED}:{seed_scope}:{g['group']}:targets".encode())
                                .hexdigest()[:16], 16))
        tokens = _per_file(g.get("tokens"), count, rng)
        rows = _per_file(g.get("rows"), count, rng)
        params = {k: v for k, v in g.items() if k not in {"group", "kind", "count", "tokens", "rows", "name"}}
        params["_seed_scope"] = seed_scope
        for i in range(count):
            name = g["name"].format(i=i, n=i + 1)
            if kind == "screenshot_image":
                fmts = g.get("formats", ["png"])
                name += "." + fmts[i % len(fmts)]
            elif kind == "sl_contract":
                fmts = g.get("formats", ["pdf"] * count)
                name += "." + fmts[i % len(fmts)]
            spec = FileSpec(scenario["id"], g["group"], i, name, kind, tokens[i], rows[i], dict(params))
            specs.append(spec)
            by_key[spec.key] = spec
    for g in groups:
        if g["kind"] != "duplicate":
            continue
        for i, src_key in enumerate(g["of"]):
            src = by_key[src_key]
            name = src.name if g.get("name", "same") == "same" else g["name"].format(i=i, n=i + 1)
            specs.append(FileSpec(scenario["id"], g["group"], i, name, src.kind, src.tokens, src.rows,
                                  dict(src.params), duplicate_of=src_key))
    return specs


_SEL = re.compile(r"^(?P<group>[A-Za-z0-9_]+)(?:\[(?P<a>-?\d*)(?P<colon>:)?(?P<b>-?\d*)\])?(?P<member>/.*)?$")


def resolve_selector(selector: str, specs: Sequence[FileSpec]) -> List[str]:
    """Resolve ``group``, ``group[i]``, ``group[a:b]`` (optionally ``/member``) to keys.

    Args:
        selector: Selector string from scenarios.yaml.
        specs: The scenario's file specs.

    Returns:
        Keys such as ``invoices[3]`` or ``zip_mixed[0]/notes/blob.bin``.
    """
    m = _SEL.match(selector)
    if not m:
        raise ValueError(f"bad selector {selector!r}")
    group = [s for s in specs if s.group == m["group"]]
    if not group:
        raise KeyError(f"unknown group in selector {selector!r}")
    if m["a"] is None and m["colon"] is None:
        chosen = group
    elif m["colon"]:
        a = int(m["a"]) if m["a"] else 0
        b = int(m["b"]) if m["b"] else len(group)
        chosen = group[a:b]
    else:
        chosen = [group[int(m["a"])]]
    keys = [s.key for s in chosen]
    member = m["member"]
    if member:
        out = []
        for s in chosen:
            names = zip_member_names(s)
            if member == "/*":
                out += [f"{s.key}/{n}" for n in names]
            else:
                out.append(f"{s.key}{member}")
        return out
    return keys


def conversations(scenario: Dict[str, Any]) -> List[List[Turn]]:
    """Expand a scenario's conversations into ``Turn`` lists (``repeat`` unrolled)."""
    specs = file_specs(scenario)
    out: List[List[Turn]] = []
    for ci, conv in enumerate(scenario.get("conversations", [])):
        turns: List[Turn] = []
        for raw in conv["turns"]:
            reps = int(raw.get("repeat", 1))
            for r in range(reps):
                attach: List[str] = []
                for sel in raw.get("attach", []) or []:
                    attach += resolve_selector(sel.replace("{i}", str(r)), specs)
                extra = {k: (v.replace("{i}", str(r)) if isinstance(v, str) else v)
                         for k, v in raw.items() if k not in {"attach", "uploading", "prompt", "repeat"}}
                turns.append(Turn(ci, len(turns), attach, int(raw.get("uploading", 0)), raw.get("prompt", ""), extra))
        out.append(turns)
    return out


def capability_matrix(scenario: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Return every capability combination a scenario should be checked under.

    Each dict holds ``variant``, ``tool_calling``, ``sandbox``, ``vision``,
    ``native_pdf``, ``window``, ``max_images``, ``model`` and, for /v1
    scenarios, ``payload`` (``text`` or ``file_parts``).
    """
    data = load_scenarios()
    model = data["models"][scenario["model"]]
    visions = scenario.get("vision_variants", [model["vision"]])
    payloads = (scenario.get("client") or {}).get("payloads") or [None]
    if scenario.get("files_from") and not scenario.get("client"):
        payloads = [None]
    out = []
    for variant in scenario["variants"]:
        caps = data["capabilities"][variant]
        for vision in visions:
            for payload in payloads:
                out.append({"variant": variant, "model": scenario["model"], "window": model["window"],
                            "native_pdf": model["native_pdf"], "max_images": model.get("max_images"),
                            "vision": vision, "payload": payload, "budget_share": data["budget_share"],
                            **caps})
    return out


def _index_matches(rule_value: Any, index: int, last: int) -> bool:
    if rule_value is None or rule_value == "*":
        return True
    if rule_value == "last":
        return index == last
    if isinstance(rule_value, list):
        return index in rule_value
    if isinstance(rule_value, str) and ":" in rule_value:
        a, b = rule_value.split(":")
        return (int(a) if a else 0) <= index < (int(b) if b else last + 1)
    return int(rule_value) == index


def expectations_for(scenario: Dict[str, Any], turn: Turn, caps: Dict[str, Any]) -> Dict[str, Any]:
    """Merge the scenario's expectation rules that apply to one turn and capability set.

    Group selectors only match files attached on that turn; ``new`` and
    ``earlier`` are the turn's own files and the conversation's earlier files.
    Later rules override earlier ones per file key and per property.

    Returns:
        ``{"files": {key: [allowed dispositions]}, "props": {...}}``.
    """
    specs = file_specs(scenario)
    convs = conversations(scenario)
    conv = convs[turn.conversation]
    new = list(dict.fromkeys(turn.attach))
    dupes = {s.key for s in specs if s.duplicate_of}
    earlier: List[str] = []
    for t in conv[: turn.index]:
        for k in t.attach:
            # A re-sent duplicate resolves to its source's ref, so it is not a separate earlier file.
            if k not in earlier and k not in new and k not in dupes:
                earlier.append(k)
    files: Dict[str, List[str]] = {}
    props: Dict[str, Any] = {}
    for rule in scenario.get("expect", []):
        if not _index_matches(rule.get("conversation"), turn.conversation, len(convs) - 1):
            continue
        if not _index_matches(rule.get("turn"), turn.index, len(conv) - 1):
            continue
        when = rule.get("when") or {}
        if any(caps.get(k) != v for k, v in when.items()):
            continue
        for sel, disp in (rule.get("files") or {}).items():
            allowed = disp if isinstance(disp, list) else [disp]
            if sel == "new":
                keys = new
            elif sel == "earlier":
                keys = earlier
            else:
                resolved = resolve_selector(sel, specs)
                keys = [k for k in resolved if k.split("/", 1)[0] in new]
            for k in keys:
                files[k] = list(allowed)
        for k, v in rule.items():
            if k not in {"conversation", "turn", "when", "files"}:
                props[k] = v
    return {"files": files, "props": props}


def disposition_matches(allowed: Iterable[str], actual: str) -> bool:
    """True when ``actual`` satisfies the allowed list (``inline`` accepts ``back_filled``)."""
    allowed = list(allowed)
    return actual in allowed or (actual == "back_filled" and "inline" in allowed)


# Fixture disposition -> planner status values (``FileStatus`` in
# docsgpt/agents/attachment_budget.py). ``image`` and ``back_filled`` are
# inline files (a native image part / inlined after a skipped file);
# ``unreadable`` has no content to send, so the planner lists it as tool or
# not_included and the manifest says why. ``skipped`` (rejected type,
# unsupported zip member) and ``duplicate`` (collapsed onto the first ref) never
# reach the planner as files of their own and map to no status.
PLANNER_STATUS: Dict[str, Tuple[str, ...]] = {
    "inline": ("inline",), "back_filled": ("inline",), "image": ("inline",), "partial": ("partial",),
    "tool": ("tool",), "sandbox": ("sandbox",), "earlier": ("earlier",), "not_included": ("not_included",),
    "unreadable": ("tool", "not_included"), "skipped": (), "duplicate": (),
}


def planner_statuses(allowed: Iterable[str]) -> set:
    """Map allowed fixture dispositions to the planner status values they permit."""
    return {status for d in allowed for status in PLANNER_STATUS[d]}


def estimated_tokens(spec: FileSpec) -> int:
    """Rough stored-text token estimate at FULL size (images 0, sheets by rows)."""
    if spec.tokens is not None:
        return int(spec.tokens)
    if spec.kind in SHEET_KINDS:
        sheets = int(spec.params.get("sheets", 1))
        per_row = {"csv_table": 22, "portfolio_xlsx": 30, "dues_xlsx": 20}.get(spec.kind, 35)
        return int((spec.rows or 0) * per_row + (sheets - 1) * 400)
    if spec.kind == "pptx_slides":
        return 120 * int(spec.params.get("slides", 8))
    if spec.kind == "mixed_zip":
        return 4000
    return 0


def attachment_records(scenario: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Virtual attachment rows at full size, for planner tests that need no files.

    ``content_hash`` is synthetic (duplicates share their source's hash);
    ``token_count`` is the full-profile target or an estimate for sheets.
    """
    out = []
    hashes: Dict[str, str] = {}
    for s in file_specs(scenario):
        src = s.duplicate_of or s.key
        if src not in hashes:
            hashes[src] = hashlib.sha256(f"{scenario['id']}:{src}".encode()).hexdigest()
        pages = s.params.get("pages")
        if isinstance(pages, list):
            pages = None
        out.append({
            "key": s.key, "group": s.group, "name": s.name, "kind": s.kind, "mime_type": mime_for(s.name),
            "token_count": estimated_tokens(s), "content_hash": hashes[src], "duplicate_of": s.duplicate_of,
            "has_text_layer": s.kind not in SCAN_KINDS | IMAGE_KINDS and s.kind != "bad_file",
            "is_image": s.kind in IMAGE_KINDS, "pages": pages, "rows": s.rows,
            "zip_members": zip_member_names(s) if s.kind.endswith("_zip") else None,
        })
    return out


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------


def _effective(value: Optional[int], cap: Optional[int]) -> Optional[int]:
    if value is None:
        return None
    return value if cap is None else min(value, cap)


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _zip_write(zf: zipfile.ZipFile, name: str, data: bytes, stored: bool = False) -> None:
    info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_TIME)
    info.compress_type = zipfile.ZIP_STORED if stored else zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    zf.writestr(info, data)


_CORE_TIME = re.compile(rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)")


def _normalize_ooxml(path: Path) -> None:
    """Rewrite an OOXML/zip file with fixed timestamps so its bytes are reproducible."""
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(path, "w") as zout:
        for name, data in items:
            if name == "docProps/core.xml":
                data = _CORE_TIME.sub(rb"\g<1>2026-01-01T00:00:00Z\g<3>", data)
            _zip_write(zout, name, data, stored=name == "mimetype")


_FONT_CACHE: Dict[str, Optional[str]] = {}
_DEVANAGARI_CANDIDATES = [
    os.environ.get("MANY_ATTACHMENTS_DEVANAGARI_FONT", ""),
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
    "/usr/share/fonts/truetype/lohit-devanagari/Lohit-Devanagari.ttf",
    "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
    "/usr/share/fonts/truetype/fonts-deva-extra/kalimati.ttf",
    "C:/Windows/Fonts/mangal.ttf",
]


def _pdf_font(script: str) -> Optional[str]:
    """Register and return a reportlab font name for ``latin`` or ``deva`` text.

    Latin text uses reportlab's bundled Vera (its text layer keeps č/š/ž).
    Devanagari needs a system TTF; returns None when none is installed, and
    callers then fall back to romanized Hindi.
    """
    if script in _FONT_CACHE:
        return _FONT_CACHE[script]
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    name: Optional[str] = None
    if script == "latin":
        pdfmetrics.registerFont(TTFont("MAVera", str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf")))
        name = "MAVera"
    else:
        for cand in _DEVANAGARI_CANDIDATES:
            if cand and Path(cand).is_file() and cand.lower().endswith(".ttf"):
                try:
                    pdfmetrics.registerFont(TTFont("MADeva", cand))
                    name = "MADeva"
                    break
                except Exception:  # noqa: BLE001 - try the next candidate
                    continue
    _FONT_CACHE[script] = name
    return name


def devanagari_available() -> bool:
    """True when a Devanagari-capable TTF was found for Hindi PDFs."""
    return _pdf_font("deva") is not None


def _write_text_pdf(path: Path, blocks: List[Tuple[str, str]], *, font: str = "latin", size: float = 9.5,
                    pages: Optional[int] = None, footer: Optional[str] = "Page {p}", figures: bool = False,
                    rng: Optional[random.Random] = None, mono: bool = False) -> int:
    """Lay out ``(heading, body)`` blocks as a text PDF; return the page count.

    Args:
        path: Output path.
        blocks: Sections to render; bodies are wrapped by character width.
        font: ``latin`` or ``deva`` (see ``_pdf_font``).
        size: Font size in points.
        pages: Spread the text over at least this many pages.
        footer: Page footer pattern (``{p}`` = page number) or None.
        figures: Draw a simple labelled diagram at the bottom of each page.
        rng: Random source for figures.
        mono: Use Courier (receipts).
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    font_name = "Courier" if mono else (_pdf_font(font) or _pdf_font("latin"))
    width, height = A4
    margin = 48
    leading = size * 1.32
    fig_h = 170 if figures else 0
    max_lines = int((height - 2 * margin - fig_h - 14) / leading)
    wrap = int((width - 2 * margin) / (size * (0.6 if mono else 0.5)))
    lines: List[Tuple[str, str]] = []
    for head, body in blocks:
        if head:
            lines.append(("H", head))
        for para in body.split("\n"):
            wrapped = textwrap.wrap(para, wrap, break_long_words=True) or [""]
            lines += [("P", ln) for ln in wrapped]
        lines.append(("P", ""))
    per_page = max_lines
    if pages:
        per_page = max(4, min(max_lines, math.ceil(len(lines) / pages)))
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    page = 0
    for start in range(0, max(1, len(lines)), per_page):
        page += 1
        y = height - margin
        for style, text in lines[start:start + per_page]:
            c.setFont(font_name, size + (2 if style == "H" else 0))
            c.drawString(margin, y, text)
            y -= leading
        if figures and rng is not None:
            _draw_figure(c, margin, margin + 18, width - 2 * margin, fig_h - 20, rng, page)
        if footer:
            c.setFont(font_name, 8)
            c.drawCentredString(width / 2, margin / 2, footer.format(p=page))
        c.showPage()
    while pages and page < pages:
        page += 1
        if footer:
            c.setFont(font_name, 8)
            c.drawCentredString(width / 2, margin / 2, footer.format(p=page))
        c.showPage()
    c.save()
    return page


def _draw_figure(c: Any, x: float, y: float, w: float, h: float, rng: random.Random, n: int) -> None:
    """Draw a cell-like diagram with labels A/B (exam-paper figure, mostly vector)."""
    c.setLineWidth(1.2)
    cx, cy = x + w / 2, y + h / 2
    c.ellipse(cx - 150, cy - 60, cx + 150, cy + 60)
    c.circle(cx - 30, cy, 22)
    for _ in range(5):
        px, py = cx + rng.randint(-120, 110), cy + rng.randint(-40, 40)
        c.ellipse(px, py, px + 26, py + 12)
    c.line(cx - 30, cy, cx + 190, cy + 40)
    c.line(cx + 60, cy - 30, cx + 190, cy - 40)
    c.setFont("Helvetica", 9)
    c.drawString(cx + 195, cy + 38, "A")
    c.drawString(cx + 195, cy - 42, "B")
    c.drawCentredString(cx, y + 2, f"Figure {n}")


def _pil_font(size: int) -> Any:
    from PIL import ImageFont

    return ImageFont.load_default(size=size)


def _render_scan_page(lines: List[str], rng: random.Random, scale: float) -> Any:
    """Render text lines as a noisy, slightly rotated grayscale 'scanned' page."""
    from PIL import Image, ImageDraw, ImageFilter

    w, h = int(1240 * scale), int(1754 * scale)
    img = Image.new("L", (w, h), 245)
    d = ImageDraw.Draw(img)
    font = _pil_font(max(10, int(26 * scale)))
    y = int(90 * scale)
    for ln in lines:
        d.text((int(90 * scale), y), ln, fill=rng.randint(10, 60), font=font)
        y += int(36 * scale)
        if y > h - 90 * scale:
            break
    for _ in range(int(400 * scale)):
        px, py = rng.randrange(w), rng.randrange(h)
        d.point((px, py), fill=rng.randint(80, 200))
    img = img.rotate(rng.uniform(-1.2, 1.2), fillcolor=245, expand=False)
    return img.filter(ImageFilter.GaussianBlur(0.6))


def _write_image_pdf(path: Path, page_images: Iterable[Any]) -> int:
    """Write an image-only PDF (no text layer) from PIL page images."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    n = 0
    for img in page_images:
        buf = io.BytesIO()
        img.convert("L").save(buf, format="JPEG", quality=55)
        buf.seek(0)
        c.drawImage(ImageReader(buf), 0, 0, A4[0], A4[1])
        c.showPage()
        n += 1
    c.save()
    return n


# ---------------------------------------------------------------------------
# Writers. Each returns a dict merged into the manifest entry; "_text" (the
# source text) is written to the sidecar and removed from the entry.
# ---------------------------------------------------------------------------


def _w_invoice_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    lines, facts = corpus.make_invoice(rng, spec.index + 1, tokens)
    _write_text_pdf(path, [("", "\n".join(lines))], mono=True, footer=None, size=9)
    return {"pages": 1, "facts": facts, "_text": "\n".join(lines)}


def _scan_lines(spec: FileSpec, rng: random.Random, page: int) -> List[str]:
    content = spec.params.get("content", "book")
    if content == "invoice":
        return corpus.make_invoice(rng, 900 + spec.index, 300)[0]
    if content == "statement":
        return _statement_lines(rng, spec.index, 40)
    if content == "land_registry":
        return _land_registry_lines(rng, page)
    return textwrap.wrap(" ".join(corpus.en_sentence(rng) for _ in range(30)), 70)[:38]


def _w_scan_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    pages = int(spec.params.get("pages", 1))
    if prof["scan_pages_cap"]:
        pages = min(pages, prof["scan_pages_cap"])
    texts: List[str] = []

    def gen():
        for p in range(pages):
            lines = _scan_lines(spec, rng, p)
            texts.append("\n".join(lines))
            yield _render_scan_page(lines, rng, max(0.5, prof["image_scale"]))

    n = _write_image_pdf(path, gen())
    ocr = corpus.ocr_noise(random.Random(spec.seed("ocr")), "\n\n".join(texts))
    return {"pages": n, "has_text_layer": False, "_text": ocr, "text_kind": "ocr"}


def _statement_lines(rng: random.Random, index: int, n: int) -> List[str]:
    lines = ["FICTIONAL COMMUNITY BANK - ACCOUNT STATEMENT", "Account holder: Example Sports Club (fictional)",
             "Account: XX00 0000 0000 0000 0000", f"Statement no. {index + 1:03d}", "",
             "Date        Description                         Debit      Credit     Balance"]
    bal = rng.randint(1000, 9000)
    for i in range(n):
        amt = rng.randint(5, 400)
        credit = rng.random() < 0.45
        bal += amt if credit else -amt
        desc = rng.choice(["Member dues M-0" + str(rng.randint(10, 99)), "Pitch hire", "Kit supplier XX-000",
                           "Tournament fee", "Bank charge", "Raffle income", "Coach expenses"])
        lines.append(f"2026-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}  {desc:<34}"
                     f"  {'' if credit else amt:>8}  {amt if credit else '':>8}  {bal:>9}")
    return lines


def _w_statement_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    lines = _statement_lines(rng, spec.index, 10)
    while corpus.count_tokens("\n".join(lines)) < tokens:
        lines += _statement_lines(rng, spec.index, 10)[6:]
    text = "\n".join(lines)
    pages = _write_text_pdf(path, [("", text)], mono=True, size=8.5)
    return {"pages": pages, "_text": text}


def _w_long_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    lang = spec.params.get("lang", "en")
    script = "deva" if devanagari_available() else "latn"
    chapters = int(spec.params.get("chapters", 4))
    if lang == "en":
        blocks = [(f"Module chapter {c + 1}", corpus.fill_to_tokens(lambda i: corpus.en_paragraph(rng),
                                                                    max(50, tokens // chapters), "\n\n"))
                  for c in range(chapters)]
    else:
        blocks = corpus.rulebook_chapters(rng, tokens, chapters, lang, script)
    font = "deva" if lang != "en" and script == "deva" else "latin"
    pages = _write_text_pdf(path, blocks, font=font)
    return {"pages": pages, "lang": lang, "script": script if lang != "en" else "latn",
            "chapters": chapters, "_text": "\n\n".join(f"{h}\n{b}" for h, b in blocks)}


def _w_exam_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    lines = corpus.exam_questions(rng, spec.index + 1, tokens)
    pr = spec.params.get("pages", [2, 6])
    pages = rng.randint(pr[0], pr[1]) if isinstance(pr, list) else int(pr)
    with_fig = spec.params.get("figures") == "half" and spec.index % 2 == 0
    n = _write_text_pdf(path, [("", "\n".join(lines))], pages=pages, figures=with_fig, rng=rng)
    return {"pages": n, "figures": with_fig, "_text": "\n".join(lines)}


def _w_report_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    years = list(range(2016, 2024))
    per = max(60, tokens // len(years))
    blocks = []
    for y in years:
        body = corpus.fill_to_tokens(lambda i, y=y: corpus.examiner_report_paragraph(rng, y, i + 1), per, "\n\n")
        blocks.append((f"Examiners' report {y} - Biology (fictional board)", body))
    n = _write_text_pdf(path, blocks)
    return {"pages": n, "_text": "\n\n".join(f"{h}\n{b}" for h, b in blocks)}


def _w_paper_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    thesis = spec.params.get("thesis_index") == spec.index
    title = (f"{'Doctoral thesis' if thesis else 'Paper'} {spec.index + 1}: "
             f"{rng.choice(['Fouling', 'Pressure drop', 'Fin geometry', 'Additive-manufactured cores'])} "
             f"in compact heat exchangers (fictional)")
    blocks = corpus.paper_sections(rng, title, tokens)
    pr = spec.params.get("thesis_pages") if thesis else spec.params.get("pages", [8, 20])
    pages = rng.randint(pr[0], pr[1]) if isinstance(pr, list) else int(pr)
    if prof["pages_cap"]:
        pages = min(pages, prof["pages_cap"])
    n = _write_text_pdf(path, blocks, pages=pages)
    return {"pages": n, "thesis": thesis, "_text": "\n\n".join(f"{h}\n{b}" for h, b in blocks)}


def _w_docx(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int,
            text: Optional[str] = None, title: Optional[str] = None) -> Dict:
    import docx

    doc = docx.Document()
    cp = doc.core_properties
    cp.created = cp.modified = cp.last_printed = FIXED_DT.replace(tzinfo=None)
    cp.author = "Fictional Author"
    if text is None:
        topic = spec.params.get("topic", "project")
        title = (f"Committee report {spec.index + 1} (fictional club)" if topic == "club"
                 else f"Project {spec.index + 1} brief (fictional coursework)")
        text = corpus.fill_to_tokens(lambda i: corpus.en_paragraph(rng, (1, 3)), max(30, tokens - 20), "\n")
    doc.add_heading(title or spec.name, level=1)
    for para in text.split("\n"):
        doc.add_paragraph(para)
    doc.save(str(path))
    _normalize_ooxml(path)
    return {"_text": f"{title}\n{text}"}


def _w_docx_writeup(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    return _w_docx(path, spec, rng, prof, tokens)


def _w_xlsx_workbook(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, rows: int) -> Dict:
    from openpyxl import Workbook

    sheets = int(spec.params.get("sheets", 3))
    formulas = bool(spec.params.get("formulas", False))
    wb = Workbook(write_only=True)
    wb.properties.created = wb.properties.modified = FIXED_DT.replace(tzinfo=None)
    ws = wb.create_sheet("raw_data")
    ws.append(["order_id", "date", "region", "product", "units", "unit_price", "revenue", "channel"])
    regions = ["North", "South", "East", "West", "Central"]
    products = ["Widget A", "Widget B", "Gadget C", "Service D", "Bundle E"]
    for r in range(2, rows + 2):
        units, price = rng.randint(1, 50), round(rng.uniform(2, 200), 2)
        ws.append([f"XX-{r:06d}", f"2025-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}", rng.choice(regions),
                   rng.choice(products), units, price, f"=E{r}*F{r}" if formulas else round(units * price, 2),
                   rng.choice(["online", "store", "partner"])])
    names = ["raw_data"]
    if sheets >= 2:
        s2 = wb.create_sheet("summary")
        s2.append(["region", "total_revenue", "orders"])
        for reg in regions:
            s2.append([reg, f'=SUMIF(raw_data!C:C,"{reg}",raw_data!G:G)' if formulas else 0,
                       f'=COUNTIF(raw_data!C:C,"{reg}")' if formulas else 0])
        names.append("summary")
    for k in range(3, sheets + 1):
        sk = wb.create_sheet(f"lookup_{k - 2}")
        sk.append(["code", "label", "weight"])
        for j in range(60):
            sk.append([f"C{j:03d}", f"Category {j}", round(rng.random(), 3)])
        names.append(f"lookup_{k - 2}")
    wb.save(str(path))
    _normalize_ooxml(path)
    return {"rows": rows, "sheets": names, "formulas": formulas}


_SEC_A = ["Contoh", "Sample", "Placeholder", "Example", "Fictive", "Teladan", "Vzorec", "Primer"]
_SEC_B = ["Industries", "Holdings", "Energy", "Bank", "Pharma", "Logistics", "Foods", "Software", "Cement", "Telecom"]
_SECTORS = ["Financials", "Energy", "Health care", "Technology", "Materials", "Consumer", "Utilities", "Industrials"]


def _w_portfolio_xlsx(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, rows: int) -> Dict:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.created = wb.properties.modified = FIXED_DT.replace(tzinfo=None)
    ws = wb.active
    ws.title = "Holdings"
    ws["A1"] = f"Monthly portfolio disclosure - Fictional Growth Fund {spec.index + 1:02d}"
    ws.merge_cells("A1:F1")
    ws["A2"] = "Portfolio as on 31-XX-2026 (fictional data)"
    ws.append([])
    ws.append(["Name of the instrument", "Code (placeholder)", "Industry / Sector", "Quantity",
               "Market value (in 000s)", "% to NAV"])
    pool = [(f"{a} {b} Ltd", f"XX{(ai * 10 + bi):010d}") for ai, a in enumerate(_SEC_A) for bi, b in enumerate(_SEC_B)]
    weights = [rng.random() ** 3 for _ in range(rows)]
    total = sum(weights)
    for (name, code), w in zip(rng.sample(pool * 6, rows), weights):
        ws.append([name, code, rng.choice(_SECTORS), rng.randint(100, 900000), round(rng.uniform(10, 9000), 2),
                   round(100 * w / total, 2)])
    ws.append([])
    ws.append(["Total", None, None, None, None, 100.0])
    ws.append(["Notes: codes are placeholders, not real identifiers."])
    wb.save(str(path))
    _normalize_ooxml(path)
    return {"rows": rows, "sheets": ["Holdings"], "header_row": 4}


def _w_dues_xlsx(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, rows: int) -> Dict:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    fills = {"paid": "C6EFCE", "partial": "FFEB9C", "unpaid": "FFC7CE"}
    wb = Workbook()
    wb.properties.created = wb.properties.modified = FIXED_DT.replace(tzinfo=None)
    ws = wb.active
    ws.title = "Dues 2026"
    ws.append(["Member", "Joined", "Dues (EUR)", "Paid on"])
    counts = {k: 0 for k in fills}
    for r in range(rows):
        status = rng.choices(list(fills), weights=[6, 2, 2])[0]
        counts[status] += 1
        ws.append([f"Member {r + 1:03d} (fictional)", f"20{rng.randint(10, 25)}", 120,
                   "" if status == "unpaid" else f"2026-0{rng.randint(1, 9)}-{rng.randint(10, 28)}"])
        for col in "ABCD":
            ws[f"{col}{r + 2}"].fill = PatternFill("solid", fgColor=fills[status])
    lg = wb.create_sheet("Legend")
    for status, colour in fills.items():
        lg.append([status, f"row colour #{colour}"])
        lg[f"A{lg.max_row}"].fill = PatternFill("solid", fgColor=colour)
    wb.save(str(path))
    _normalize_ooxml(path)
    return {"rows": rows, "sheets": ["Dues 2026", "Legend"], "facts": {"status_by_colour": counts}}


def _w_csv_table(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, rows: int) -> Dict:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["id", "date", "category", "value", "flag", "comment"])
        for r in range(rows):
            w.writerow([r + 1, f"2025-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}",
                        rng.choice(["alpha", "beta", "gamma", "delta"]), round(rng.gauss(100, 25), 3),
                        rng.choice(["", "", "check", "outlier"]), rng.choice(["", "", "late entry", "manual fix"])])
    return {"rows": rows}


def _w_notebook_html(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    def cell(i: int) -> str:
        if i % 3 == 0:
            return f"[{i}] import pandas as pd\ndf = pd.read_csv('raw_data_{spec.index + 1}.csv')\ndf.describe()"
        if i % 3 == 1:
            rows = " | ".join(f"{rng.choice(['alpha', 'beta', 'gamma'])} {round(rng.gauss(100, 20), 2)}"
                              for _ in range(8))
            return f"Out[{i}]: {rows}"
        return corpus.en_paragraph(rng, (2, 4))

    text = corpus.fill_to_tokens(cell, tokens, "\n\n")
    body = "".join(f'<div class="cell"><pre>{c}</pre></div>\n' for c in text.split("\n\n"))
    html = (f"<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>analysis_notebook_{spec.index + 1}.ipynb"
            f" - Colaboratory (fictional export)</title></head><body>\n{body}</body></html>\n")
    path.write_text(html, encoding="utf-8")
    return {"_text": text}


def _w_pptx_slides(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    from pptx import Presentation

    prs = Presentation()
    prs.core_properties.created = prs.core_properties.modified = FIXED_DT.replace(tzinfo=None)
    n = int(spec.params.get("slides", 8))
    for i in range(n):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Finding {i + 1} (fictional)"
        s.placeholders[1].text = "\n".join(corpus.en_sentence(rng) for _ in range(3))
    prs.save(str(path))
    _normalize_ooxml(path)
    return {"slides": n}


def _w_log_txt(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    if spec.params.get("style") == "transcript":
        text = corpus.fill_to_tokens(lambda i: f"{'Lecturer' if i % 4 else 'Student'}: {corpus.en_sentence(rng)}",
                                     tokens)
    else:
        text = corpus.fill_to_tokens(lambda i: corpus.log_line(rng, i), tokens)
    path.write_text(text, encoding="utf-8")
    return {"_text": text}


def _w_python_script(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    def unit(i: int) -> str:
        return (f"def parse_block_{i}(line: str) -> dict:\n    \"\"\"{corpus.en_sentence(rng)}\"\"\"\n"
                f"    parts = line.split(' ', {rng.randint(2, 6)})\n    return {{'level': parts[1], 'raw': line}}\n")

    text = "import re\nimport sys\n\n" + corpus.fill_to_tokens(unit, tokens, "\n")
    path.write_text(text, encoding="utf-8")
    return {"_text": text}


def _w_xml_dump(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    body = corpus.fill_to_tokens(lambda i: corpus.xml_config_line(rng, i), tokens)
    text = f"<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n<config vendor=\"Fictional Networks\">\n{body}\n</config>\n"
    path.write_text(text, encoding="utf-8")
    return {"_text": text}


def _w_epub_book(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    chapters = 10
    per = max(50, tokens // chapters)
    texts = [corpus.fill_to_tokens(lambda i: corpus.en_paragraph(rng), per, "\n\n") for _ in range(chapters)]
    with zipfile.ZipFile(path, "w") as zf:
        _zip_write(zf, "mimetype", b"application/epub+zip", stored=True)
        _zip_write(zf, "META-INF/container.xml", (
            '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>").encode())
        manifest = "".join(f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>'
                           for i in range(chapters))
        spine = "".join(f'<itemref idref="c{i}"/>' for i in range(chapters))
        _zip_write(zf, "OEBPS/content.opf", (
            '<?xml version="1.0"?><package version="3.0" xmlns="http://www.idpf.org/2007/opf" unique-identifier="id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:uuid:00000000-0000-'
            f'0000-0000-00000000000{spec.index}</dc:identifier><dc:title>Course book {spec.index + 1} (fictional)'
            f"</dc:title><dc:language>en</dc:language></metadata><manifest>{manifest}</manifest>"
            f"<spine>{spine}</spine></package>").encode())
        for i, t in enumerate(texts):
            paras = "".join(f"<p>{p}</p>" for p in t.split("\n\n"))
            _zip_write(zf, f"OEBPS/c{i}.xhtml", (
                '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml"><head><title>'
                f"Chapter {i + 1}</title></head><body><h1>Chapter {i + 1}</h1>{paras}</body></html>").encode())
    return {"chapters": chapters, "_text": "\n\n".join(f"Chapter {i + 1}\n{t}" for i, t in enumerate(texts))}


# --- images ----------------------------------------------------------------


def _save_image(img: Any, path: Path) -> None:
    ext = path.suffix.lower()
    if ext in (".jpg", ".jpeg"):
        img.convert("RGB").save(path, format="JPEG", quality=82)
    elif ext == ".webp":
        img.convert("RGB").save(path, format="WEBP", quality=80, method=4)
    else:
        img.save(path, format="PNG", optimize=False)


def _scaled(spec: FileSpec, prof: Dict) -> Tuple[int, int]:
    w, h = spec.params.get("size", [1080, 2400])
    s = prof["image_scale"]
    return max(200, int(w * s)), max(200, int(h * s))


def _w_screenshot_image(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    from PIL import Image, ImageDraw

    w, h = _scaled(spec, prof)
    img = Image.new("RGB", (w, h), (236, 229, 221))
    d = ImageDraw.Draw(img)
    unit = w / 1080
    d.rectangle([0, 0, w, int(150 * unit)], fill=(18, 140, 126))
    d.text((int(40 * unit), int(55 * unit)), "Reading group (fictional)", fill="white", font=_pil_font(int(44 * unit)))
    lo, hi = spec.params.get("quotes", [1, 3])
    quotes = [corpus.make_quote(rng) for _ in range(rng.randint(lo, hi))]
    y = int(220 * unit)
    font = _pil_font(max(12, int(40 * unit)))
    for q in quotes:
        lines = textwrap.wrap(f"\u201c{q['text']}\u201d - {q['author']}", 34)
        bh = int((len(lines) * 54 + 50) * unit)
        x0 = int(rng.choice([40, 180]) * unit)
        d.rounded_rectangle([x0, y, x0 + int(820 * unit), y + bh], radius=int(24 * unit), fill=(255, 255, 255))
        for k, ln in enumerate(lines):
            d.text((x0 + int(30 * unit), y + int((25 + k * 54) * unit)), ln, fill=(20, 20, 20), font=font)
        y += bh + int(60 * unit)
    _save_image(img, path)
    return {"width": w, "height": h, "facts": {"quotes": quotes}}


def _w_photo_image(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    from PIL import Image, ImageDraw, ImageFilter

    w, h = _scaled(spec, prof)
    subject = spec.params.get("subject") or corpus.PHOTO_SUBJECTS[spec.index % len(corpus.PHOTO_SUBJECTS)]
    img = Image.new("RGB", (w, h), tuple(rng.randint(90, 200) for _ in range(3)))
    d = ImageDraw.Draw(img)
    for _ in range(12):
        x, y = rng.randrange(w), rng.randrange(h)
        d.rectangle([x, y, x + rng.randint(20, w // 3), y + rng.randint(20, h // 3)],
                    fill=tuple(rng.randint(30, 230) for _ in range(3)))
    bx, by = w // 6, h // 5
    d.rectangle([bx, by, w - bx, h - by], fill=(250, 248, 240), outline=(60, 60, 60), width=4)
    font = _pil_font(max(12, w // 40))
    if subject == "receipt":
        lines = ["FICTIONAL SPORTS SHOP", "Receipt no. XX-000", "Training cones x 20   EUR 34.00",
                 "Bibs x 12            EUR 48.00", "TOTAL                EUR 82.00", "Thank you"]
    else:
        lines = [subject.upper() + " (fictional)"] + textwrap.wrap(corpus.en_paragraph(rng, (2, 3)), 38)[:6]
    for k, ln in enumerate(lines):
        d.text((bx + 20, by + 20 + k * (w // 30)), ln, fill=(25, 25, 25), font=font)
    if spec.index % 5 == 1:
        img = img.rotate(90, expand=True)
    if spec.index % 7 == 3:
        img = img.filter(ImageFilter.GaussianBlur(3))
    _save_image(img, path)
    return {"width": img.width, "height": img.height, "facts": {"subject": subject, "lines": lines}}


def _w_quiz_image(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    from PIL import Image, ImageDraw

    w, h = _scaled(spec, prof)
    quiz = corpus.make_quiz(rng, spec.index)
    img = Image.new("RGB", (w, h), (250, 250, 252))
    d = ImageDraw.Draw(img)
    font = _pil_font(max(12, h // 22))
    d.text((w // 20, h // 10), quiz["question"], fill=(10, 10, 10), font=font)
    for k, opt in enumerate(quiz["options"]):
        y = h // 4 + k * h // 7
        d.rounded_rectangle([w // 20, y, w - w // 20, y + h // 9], radius=10, outline=(120, 120, 140), width=2)
        d.text((w // 20 + 20, y + h // 40), f"{'ABCD'[k]}. {opt}", fill=(30, 30, 30), font=font)
    _save_image(img, path)
    return {"width": w, "height": h, "facts": quiz}


def _w_map_image(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    from PIL import Image, ImageDraw

    w, h = _scaled(spec, prof)
    colours = {"SSm": (250, 210, 120), "SK": (230, 170, 90), "K1": (190, 230, 140), "G": (120, 180, 110),
               "IG": (190, 160, 210)}
    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    zones = []
    for _ in range(14):
        cx, cy, r = rng.randrange(w * 3 // 4), rng.randrange(h), rng.randint(w // 16, w // 6)
        pts = [(cx + r * math.cos(a) * rng.uniform(0.6, 1.2), cy + r * math.sin(a) * rng.uniform(0.6, 1.2))
               for a in [k * math.pi / 4 for k in range(8)]]
        code = rng.choice(list(colours))
        zones.append(code)
        d.polygon(pts, fill=colours[code], outline=(80, 80, 80))
    d.polygon([(w * 0.3, h * 0.4), (w * 0.38, h * 0.38), (w * 0.4, h * 0.5), (w * 0.31, h * 0.52)],
              outline=(200, 0, 0), width=3)
    font = _pil_font(max(10, h // 45))
    d.text((w * 0.33, h * 0.44), "A", fill=(200, 0, 0), font=font)
    d.text((w * 0.36, h * 0.47), "B", fill=(200, 0, 0), font=font)
    lx = int(w * 0.8)
    for k, (code, col) in enumerate(colours.items()):
        d.rectangle([lx, 40 + k * 40, lx + 30, 64 + k * 40], fill=col, outline=(0, 0, 0))
        d.text((lx + 40, 42 + k * 40), code, fill=(0, 0, 0), font=font)
    _save_image(img, path)
    return {"width": w, "height": h, "facts": {"zones": sorted(set(zones))}}


def _w_bad_file(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    if spec.params.get("mode") == "corrupt_image":
        from PIL import Image

        buf = io.BytesIO()
        Image.new("RGB", (64, 64), (200, 10, 10)).save(buf, format="PNG")
        path.write_bytes(buf.getvalue()[:60])
        return {"reason": "truncated PNG"}
    path.write_bytes(b"\x00\x00\x00\x18ftypheic" + bytes(rng.randrange(256) for _ in range(2048)))
    return {"reason": "unsupported type .heic"}


# --- Slovenian legal set ---------------------------------------------------


def _land_registry_lines(rng: random.Random, page: int) -> List[str]:
    return ["IZPISEK IZ ZEMLJIŠKE KNJIGE (FIKTIVNI)", corpus.SL_HEADER, f"Stran {page + 1}",
            f"Nepremičnina: {corpus.SL_PARCELS[page % len(corpus.SL_PARCELS)]}",
            "ID znak: parcela 0000 0000/0", "Lastnik: Janez Primer, Vzorčna ulica 0, 0000 Primerjevo",
            "Delež: 1/1", "Plombe: ni", "Zaznambe: ni vpisanih zaznamb",
            "Služnosti: služnost hoje in vožnje v korist parc. št. 0000/2 (fiktivno)"] + \
        textwrap.wrap(" ".join(corpus.sl_article(rng, 1).split("\n")[2:]), 70)


def _w_sl_ordinance_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    text = corpus.sl_ordinance(rng, spec.params.get("title", "Odlok (fiktivni)"), tokens)
    n = _write_text_pdf(path, [("", text)], footer="Stran {p}")
    return {"pages": n, "lang": "sl", "_text": text}


def _w_sl_contract(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    kinds = spec.params.get("contract_kinds", ["prodajna"])
    text = corpus.sl_contract(rng, kinds[spec.index % len(kinds)], tokens)
    if path.suffix.lower() == ".docx":
        title, body = text.split("\n\n", 2)[1], text
        info = _w_docx(path, spec, rng, prof, tokens, text=body, title=title)
        return {"lang": "sl", **info}
    n = _write_text_pdf(path, [("", text)], footer="Stran {p}")
    return {"pages": n, "lang": "sl", "_text": text}


def _w_sl_certificate_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    text = corpus.sl_location_certificate(rng, tokens)
    n = _write_text_pdf(path, [("", text)], footer="Stran {p}")
    return {"pages": n, "lang": "sl", "_text": text}


def _w_sl_scan_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    spec.params.setdefault("content", "land_registry")
    info = _w_scan_pdf(path, spec, rng, {**prof, "scan_pages_cap": None}, tokens)
    return {**info, "lang": "sl"}


def _w_sl_filing_pdf(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    kind = ["TOŽBA (FIKTIVNA)", "IZVEDENSKO MNENJE (FIKTIVNO)", "ODGOVOR NA TOŽBO (FIKTIVNI)",
            "ZAPISNIK O OGLEDU (FIKTIVNI)"][spec.index % 4]
    head = (f"{corpus.SL_HEADER}\n\n{kind}\nOpr. št. XX-000/2026\nTožeča stranka: {corpus.SL_PARTIES['seller']}\n"
            f"Tožena stranka: {corpus.SL_PARTIES['buyer']}\n")
    claims = ["Tožeča stranka predlaga zaslišanje priče Marije Zgled.",
              "Izvedenec ugotavlja, da je vrednost nepremičnine na dan ogleda 000.000 EUR.",
              "Iz listinskih dokazov izhaja, da je bila kupnina plačana le delno.",
              "Tožena stranka ugovarja aktivni legitimaciji tožeče stranke.",
              "Sodišče naj izvede dokaz z ogledom na kraju samem."]
    body = corpus.fill_to_tokens(lambda i: f"{i + 1}. " + " ".join(rng.choice(claims) for _ in range(3)),
                                 max(50, tokens - corpus.count_tokens(head)), "\n\n")
    pr = spec.params.get("pages", [3, 20])
    pages = rng.randint(pr[0], pr[1])
    if prof["pages_cap"]:
        pages = min(pages, prof["pages_cap"])
    n = _write_text_pdf(path, [("", head + "\n" + body)], footer="Stran {p}", pages=pages)
    return {"pages": n, "lang": "sl", "_text": head + "\n" + body}


def _w_sl_extract_text(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    text = corpus.sl_ordinance(rng, f"Izvleček {spec.index + 1}: Zakon o vzorčnih razmerjih (fiktivni)", tokens)
    path.write_text(text, encoding="utf-8")
    return {"lang": "sl", "_text": text}


# --- archives --------------------------------------------------------------

_NEXT_DIRS = ["components", "lib", "app/(shop)", "app/api", "hooks"]


def zip_member_names(spec: FileSpec) -> List[str]:
    """Leaf member paths of an archive spec (nested zips expanded as ``inner.zip/<name>``)."""
    if spec.kind == "nextjs_zip":
        n = int(spec.params.get("members", 40))
        fixed = ["package.json", "next.config.js", "tsconfig.json", "README.md", "app/layout.tsx", "app/page.tsx",
                 "styles/globals.css", "public/logo.png"]
        rest = [f"{_NEXT_DIRS[i % len(_NEXT_DIRS)]}/Part{i:02d}.tsx" for i in range(n - len(fixed))]
        return fixed + rest
    if spec.kind == "markdown_zip":
        return [f"docs/{i + 1:02d}-topic.md" for i in range(int(spec.params.get("members", 25)))]
    if spec.kind == "mixed_zip":
        return ["README.txt", "notes/meeting.md", "data/members.csv", "images/diagram.png", "docs/summary.pdf",
                "notes/blob.bin", "inner.zip/inner_a.txt", "inner.zip/inner_b.md"]
    return []


def _png_bytes(rng: random.Random, w: int = 320, h: int = 200, label: str = "diagram") -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(img)
    for k in range(4):
        d.rectangle([20 + k * 70, 60, 70 + k * 70, 140], outline=(0, 0, 0), width=2)
    d.text((20, 20), label, fill=(0, 0, 0), font=_pil_font(16))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _w_nextjs_zip(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    names = zip_member_names(spec)
    text_members = [n for n in names if not n.endswith(".png")]
    per = max(30, tokens // len(text_members))
    texts: Dict[str, str] = {}
    with zipfile.ZipFile(path, "w") as zf:
        for name in names:
            if name.endswith(".png"):
                _zip_write(zf, name, _png_bytes(rng, label="logo"))
                continue
            if name == "package.json":
                body = json.dumps({"name": "storefront-fictional", "private": True,
                                   "scripts": {"dev": "next dev", "build": "next build"},
                                   "dependencies": {"next": "0.0.0-fixture", "react": "0.0.0-fixture"}}, indent=2)
            elif name.endswith(".tsx"):
                body = corpus.ts_component(rng, Path(name).stem.replace("(", "").replace(")", ""), per)
            elif name.endswith(".md"):
                body = corpus.markdown_doc(rng, "Storefront (fictional)", per)
            else:
                body = "\n".join(f"/* {corpus.en_sentence(rng)} */" for _ in range(max(1, per // 25)))
            texts[name] = body
            _zip_write(zf, name, body.encode())
    return {"members": names, "_text": "\n\n".join(f"// {k}\n{v}" for k, v in texts.items())}


def _w_markdown_zip(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: int) -> Dict:
    names = zip_member_names(spec)
    per = max(30, tokens // len(names))
    texts = {}
    with zipfile.ZipFile(path, "w") as zf:
        for i, name in enumerate(names):
            texts[name] = corpus.markdown_doc(rng, f"Product guide part {i + 1} (fictional)", per)
            _zip_write(zf, name, texts[name].encode())
    return {"members": names, "_text": "\n\n".join(texts.values())}


def _w_mixed_zip(path: Path, spec: FileSpec, rng: random.Random, prof: Dict, tokens: Optional[int]) -> Dict:
    pdf_tmp = path.with_suffix(".member.pdf")
    _write_text_pdf(pdf_tmp, [("Summary (fictional)", corpus.en_paragraph(rng, (8, 12)))])
    pdf_bytes = pdf_tmp.read_bytes()
    pdf_tmp.unlink()
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as iz:
        _zip_write(iz, "inner_a.txt", ("Inner note A (fictional). " + corpus.en_paragraph(rng)).encode())
        _zip_write(iz, "inner_b.md", ("# Inner B\n\n" + corpus.en_paragraph(rng)).encode())
    rows = "\n".join(f"Member {i:03d},2026,{rng.choice(['paid', 'unpaid'])}" for i in range(40))
    with zipfile.ZipFile(path, "w") as zf:
        _zip_write(zf, "README.txt", ("Handover bundle (fictional). " + corpus.en_paragraph(rng)).encode())
        _zip_write(zf, "notes/meeting.md", corpus.markdown_doc(rng, "Meeting notes (fictional)", 400).encode())
        _zip_write(zf, "data/members.csv", ("member,year,status\n" + rows).encode())
        _zip_write(zf, "images/diagram.png", _png_bytes(rng))
        _zip_write(zf, "docs/summary.pdf", pdf_bytes)
        _zip_write(zf, "notes/blob.bin", bytes(rng.randrange(256) for _ in range(4096)))
        _zip_write(zf, "inner.zip", inner.getvalue())
    return {"members": zip_member_names(spec), "unsupported_members": ["notes/blob.bin"],
            "nested_archives": ["inner.zip"]}


WRITERS = {
    "invoice_pdf": _w_invoice_pdf, "scan_pdf": _w_scan_pdf, "statement_pdf": _w_statement_pdf,
    "long_pdf": _w_long_pdf, "exam_pdf": _w_exam_pdf, "report_pdf": _w_report_pdf, "paper_pdf": _w_paper_pdf,
    "docx_writeup": _w_docx_writeup, "xlsx_workbook": _w_xlsx_workbook, "portfolio_xlsx": _w_portfolio_xlsx,
    "dues_xlsx": _w_dues_xlsx, "csv_table": _w_csv_table, "notebook_html": _w_notebook_html,
    "pptx_slides": _w_pptx_slides, "log_txt": _w_log_txt, "python_script": _w_python_script,
    "xml_dump": _w_xml_dump, "epub_book": _w_epub_book, "screenshot_image": _w_screenshot_image,
    "photo_image": _w_photo_image, "quiz_image": _w_quiz_image, "map_image": _w_map_image, "bad_file": _w_bad_file,
    "sl_ordinance_pdf": _w_sl_ordinance_pdf, "sl_contract": _w_sl_contract,
    "sl_certificate_pdf": _w_sl_certificate_pdf, "sl_scan_pdf": _w_sl_scan_pdf, "sl_filing_pdf": _w_sl_filing_pdf,
    "sl_extract_text": _w_sl_extract_text, "nextjs_zip": _w_nextjs_zip, "markdown_zip": _w_markdown_zip,
    "mixed_zip": _w_mixed_zip,
}


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------


def generate_file(spec: FileSpec, scenario_dir: Path, profile: str = "small",
                  sources: Optional[Dict[str, Dict[str, Any]]] = None) -> Dict[str, Any]:
    """Write one file and return its manifest entry.

    Args:
        spec: The file to write.
        scenario_dir: Output directory of the scenario.
        profile: ``small`` or ``full``.
        sources: Already-written entries by key (needed for duplicates).
    """
    prof = PROFILES[profile]
    out = scenario_dir / spec.rel_path
    out.parent.mkdir(parents=True, exist_ok=True)
    entry: Dict[str, Any] = {"key": spec.key, "group": spec.group, "index": spec.index, "name": spec.name,
                             "kind": spec.kind, "path": spec.rel_path, "mime": mime_for(spec.name),
                             "full_target_tokens": spec.tokens, "full_rows": spec.rows}
    if spec.duplicate_of:
        src = (sources or {})[spec.duplicate_of]
        shutil.copyfile(scenario_dir / src["path"], out)
        entry.update({k: v for k, v in src.items() if k not in entry and k not in {"sha256", "bytes"}})
        entry["duplicate_of"] = spec.duplicate_of
    else:
        rng = random.Random(spec.seed())
        if spec.kind in SHEET_KINDS:
            arg = _effective(spec.rows, prof["rows_cap"])
            entry["target_rows"] = arg
        else:
            arg = _effective(spec.tokens, prof["token_cap"])
            if spec.kind in TEXT_KINDS:
                entry["target_tokens"] = arg
        info = WRITERS[spec.kind](out, spec, rng, prof, arg)
        text = info.pop("_text", None)
        entry.setdefault("has_text_layer", spec.kind not in SCAN_KINDS | IMAGE_KINDS | {"bad_file"})
        entry.update(info)
        if text is not None:
            tp = scenario_dir / "_text" / (spec.rel_path + ".txt")
            tp.parent.mkdir(parents=True, exist_ok=True)
            tp.write_text(text, encoding="utf-8")
            entry["text_path"] = str(tp.relative_to(scenario_dir))
    entry["bytes"] = out.stat().st_size
    entry["sha256"] = _sha256(out)
    return entry


def generate_scenario(scenario_id: str, out_dir: Path, profile: str = "small",
                      groups: Optional[Sequence[str]] = None) -> Dict[str, Any]:
    """Generate every file of one scenario into ``out_dir/<scenario_id>``.

    Args:
        scenario_id: Scenario id, e.g. ``RC-04``.
        out_dir: Root output directory.
        profile: ``small`` (fast, capped sizes) or ``full`` (production-shaped sizes).
        groups: Restrict to these file groups (duplicates pull in their sources).

    Returns:
        The scenario manifest (also written to ``manifest.json``).
    """
    scenario = get_scenario(scenario_id)
    sdir = Path(out_dir) / scenario_id
    sdir.mkdir(parents=True, exist_ok=True)
    specs = file_specs(scenario)
    if groups:
        wanted = set(groups)
        needed = {s.duplicate_of for s in specs if s.group in wanted and s.duplicate_of}
        specs = [s for s in specs if s.group in wanted or s.key in needed]
    entries: Dict[str, Dict[str, Any]] = {}
    for spec in specs:
        entries[spec.key] = generate_file(spec, sdir, profile, entries)
    manifest = {
        "scenario": scenario_id, "title": scenario.get("title"), "profile": profile, "model": scenario["model"],
        "devanagari_font": devanagari_available(), "files": list(entries.values()),
        "conversations": [[{"attach": t.attach, "uploading": t.uploading, "prompt": t.prompt, **t.extra}
                           for t in conv] for conv in conversations(scenario)],
    }
    (sdir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    return manifest


def generate(out_dir: Path, scenario_list: Optional[Sequence[str]] = None,
             profile: str = "small") -> Dict[str, Dict[str, Any]]:
    """Generate several scenarios (all by default); return manifests by id."""
    return {sid: generate_scenario(sid, Path(out_dir), profile) for sid in (scenario_list or scenario_ids())}


def load_manifest(out_dir: Path, scenario_id: str) -> Dict[str, Any]:
    """Read a previously generated ``manifest.json``."""
    return json.loads((Path(out_dir) / scenario_id / "manifest.json").read_text(encoding="utf-8"))


def main(argv: Optional[Sequence[str]] = None) -> int:
    """CLI entry point."""
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("scenarios", nargs="*", help="scenario ids (default: all)")
    ap.add_argument("--out", required=True, type=Path, help="output directory")
    ap.add_argument("--full", action="store_true", help="production-shaped sizes (slow, large)")
    ap.add_argument("--list", action="store_true", help="list scenario ids and exit")
    args = ap.parse_args(argv)
    if args.list:
        for s in load_scenarios()["scenarios"]:
            print(f"{s['id']:6} {s['title']}")
        return 0
    manifests = generate(args.out, args.scenarios or None, "full" if args.full else "small")
    for sid, m in manifests.items():
        total = sum(f["bytes"] for f in m["files"])
        print(f"{sid}: {len(m['files'])} files, {total / 1e6:.1f} MB -> {args.out / sid}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
