"""The docx (``document``) artifact renderer and its spec schema.

The renderer runs in the sandbox; these tests run the same fixed program
in-process against a temp dir and read the result back with python-docx.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from docsgpt.agents.tools.artifact_docx import build_document, clean_text, render_docx_spec  # noqa: E402
from docsgpt.agents.tools.artifact_generator import (  # noqa: E402
    _RENDERERS,
    ArtifactGeneratorTool,
    _apply_spec_append,
)

# The document renderer as it was before blocks existed. Legacy specs must
# come out of the new renderer with the same body.
_LEGACY_PROGRAM = (
    "import json\n"
    "from docx import Document\n"
    "spec = json.load(open({spec_path!r}))\n"
    "doc = Document()\n"
    "title = spec.get('title')\n"
    "if title:\n"
    "    doc.add_heading(str(title), level=0)\n"
    "for sec in spec.get('sections', []):\n"
    "    heading = sec.get('heading')\n"
    "    if heading:\n"
    "        doc.add_heading(str(heading), level=1)\n"
    "    for p in (sec.get('paragraphs') or []):\n"
    "        doc.add_paragraph(str(p))\n"
    "doc.save({out_path!r})\n"
)


def _run_program(program: str, spec: dict, tmp_path: Path, name: str = "out") -> Path:
    spec_path = tmp_path / f"{name}.json"
    out_path = tmp_path / f"{name}.docx"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    namespace: dict = {}
    exec(compile(program.format(spec_path=str(spec_path), out_path=str(out_path)), "<r>", "exec"), namespace)  # noqa: S102
    return out_path


def _render(spec: dict, tmp_path: Path):
    """Render with the module the sandbox runs, imported here so its lines are measured."""
    spec_path = tmp_path / "direct.json"
    out_path = tmp_path / "direct.docx"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    render_docx_spec(str(spec_path), str(out_path))
    return Document(str(out_path))


def _render_in_sandbox_program(spec: dict, tmp_path: Path):
    """Render through the fixed program ``artifact_generator`` sends to the sandbox."""
    return Document(str(_run_program(_RENDERERS["document"], spec, tmp_path)))


def _tool() -> ArtifactGeneratorTool:
    return ArtifactGeneratorTool(tool_config={"conversation_id": "conv-1"}, user_id="u-1")


def _num_pr(paragraph):
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if num_pr is None:
        return None
    return int(num_pr.numId.val), int(num_pr.ilvl.val)


def _abstract_for(doc, num_id: int):
    numbering = doc.part.numbering_part.element
    num = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == num_id)
    abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    return next(
        a for a in numbering.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id
    ), num


def _level(abstract, ilvl: int):
    lvl = next(lv for lv in abstract.findall(qn("w:lvl")) if lv.get(qn("w:ilvl")) == str(ilvl))
    return lvl.find(qn("w:numFmt")).get(qn("w:val")), lvl.find(qn("w:lvlText")).get(qn("w:val"))


# ---------------------------------------------------------------------------
# Backwards compatibility
# ---------------------------------------------------------------------------


def test_legacy_sections_spec_renders_the_same_body(tmp_path):
    spec = {
        "title": "Quarterly report",
        "sections": [
            {"heading": "Summary", "paragraphs": ["Revenue rose.", "Costs * fell _a bit_."]},
            {"paragraphs": ["No heading here."]},
            {"heading": "Empty"},
        ],
    }
    old = Document(str(_run_program(_LEGACY_PROGRAM, spec, tmp_path, "old")))
    new = Document(str(_run_program(_RENDERERS["document"], spec, tmp_path, "new")))
    assert new.element.body.xml == old.element.body.xml


def test_legacy_paragraph_markup_stays_literal(tmp_path):
    doc = _render({"sections": [{"paragraphs": ["**not bold** [x](https://e.com)"]}]}, tmp_path)
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "**not bold** [x](https://e.com)"
    assert all(not run.bold for run in paragraph.runs)


def test_sections_and_blocks_render_in_order(tmp_path):
    spec = {
        "sections": [{"heading": "Old", "paragraphs": ["old text"]}],
        "blocks": [{"type": "paragraph", "text": "new text"}],
    }
    doc = _render(spec, tmp_path)
    assert [p.text for p in doc.paragraphs] == ["Old", "old text", "new text"]


# ---------------------------------------------------------------------------
# Blocks
# ---------------------------------------------------------------------------


def test_headings_use_word_heading_styles(tmp_path):
    doc = _render(
        {
            "title": "Agreement",
            "blocks": [
                {"type": "heading", "text": "Definitions"},
                {"type": "heading", "text": "Interpretation", "level": 3, "align": "center"},
                {"type": "heading", "runs": [{"text": "Fees", "italic": True}], "level": 6},
            ],
        },
        tmp_path,
    )
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert styles == [
        ("Title", "Agreement"),
        ("Heading 1", "Definitions"),
        ("Heading 3", "Interpretation"),
        ("Heading 6", "Fees"),
    ]
    assert doc.paragraphs[2].alignment == 1  # WD_ALIGN_PARAGRAPH.CENTER
    assert doc.paragraphs[3].runs[0].italic is True


def test_runs_carry_inline_formatting(tmp_path):
    doc = _render(
        {
            "blocks": [
                {
                    "type": "paragraph",
                    "runs": [
                        "The ",
                        {"text": "Supplier", "bold": True},
                        " shall ",
                        {"text": "not", "italic": True, "underline": True},
                        " ",
                        {"text": "deleted", "strike": True},
                        " ",
                        {"text": "clause_7", "code": True},
                    ],
                    "align": "justify",
                }
            ]
        },
        tmp_path,
    )
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "The Supplier shall not deleted clause_7"
    runs = {run.text: run for run in paragraph.runs}
    assert runs["Supplier"].bold is True
    assert runs["The "].bold is None
    assert runs["not"].italic is True and runs["not"].underline is True
    assert runs["deleted"].font.strike is True
    assert runs["clause_7"].font.name == "Courier New"
    assert paragraph.alignment == 3  # WD_ALIGN_PARAGRAPH.JUSTIFY


def test_link_runs_become_external_hyperlinks(tmp_path):
    doc = _render(
        {
            "blocks": [
                {
                    "type": "paragraph",
                    "runs": [
                        "See ",
                        {"text": "the register", "link": "https://example.com/register?a=1&b=2", "bold": True},
                        " or ",
                        {"text": "write", "link": "mailto:legal@example.com"},
                    ],
                }
            ]
        },
        tmp_path,
    )
    p = doc.paragraphs[0]._p
    links = p.findall(qn("w:hyperlink"))
    assert len(links) == 2
    targets = {doc.part.rels[link.get(qn("r:id"))].target_ref for link in links}
    assert targets == {"https://example.com/register?a=1&b=2", "mailto:legal@example.com"}
    first_run = links[0].find(qn("w:r"))
    assert first_run.find(qn("w:rPr")).find(qn("w:b")) is not None
    assert "".join(t.text for t in links[0].iter(qn("w:t"))) == "the register"


def test_renderer_draws_an_unsafe_link_as_plain_text():
    # The schema refuses it; the renderer still never writes it as a link.
    doc = build_document({"blocks": [{"type": "paragraph", "runs": [{"text": "x", "link": "javascript:alert(1)"}]}]})
    assert doc.paragraphs[0]._p.findall(qn("w:hyperlink")) == []
    assert doc.paragraphs[0].text == "x"


def test_ordered_lists_number_clauses_as_an_outline(tmp_path):
    doc = _render(
        {
            "blocks": [
                {
                    "type": "list",
                    "ordered": True,
                    "items": [
                        "Definitions",
                        {
                            "text": "Services",
                            "items": ["Scope", {"runs": [{"text": "Delivery", "bold": True}], "items": ["Deadline"]}],
                        },
                        "Fees",
                    ],
                }
            ]
        },
        tmp_path,
    )
    numbered = [(p.text, _num_pr(p)) for p in doc.paragraphs]
    num_id = numbered[0][1][0]
    assert [(text, lvl) for text, (nid, lvl) in numbered] == [
        ("Definitions", 0),
        ("Services", 0),
        ("Scope", 1),
        ("Delivery", 1),
        ("Deadline", 2),
        ("Fees", 0),
    ]
    assert {nid for _, (nid, _) in numbered} == {num_id}
    abstract, _ = _abstract_for(doc, num_id)
    assert _level(abstract, 0) == ("decimal", "%1.")
    assert _level(abstract, 1) == ("decimal", "%1.%2.")
    assert _level(abstract, 2) == ("decimal", "%1.%2.%3.")
    assert all(p.style.name == "List Paragraph" for p in doc.paragraphs)
    assert doc.paragraphs[3].runs[0].bold is True


def test_each_list_restarts_and_start_overrides_the_first_number(tmp_path):
    doc = _render(
        {
            "blocks": [
                {"type": "list", "ordered": True, "items": ["a", "b"]},
                {"type": "list", "ordered": True, "items": ["c"], "start": 5},
                {"type": "list", "items": ["x", {"text": "y", "items": ["z"]}]},
            ]
        },
        tmp_path,
    )
    ids = [_num_pr(p)[0] for p in doc.paragraphs]
    first, second, bullets = ids[0], ids[2], ids[3]
    assert ids[1] == first and len({first, second, bullets}) == 3
    abstract_1, _ = _abstract_for(doc, first)
    abstract_2, num_2 = _abstract_for(doc, second)
    # Both numbered lists share one definition; the second restarts at 5.
    assert abstract_1 is not None and abstract_1.get(qn("w:abstractNumId")) == abstract_2.get(qn("w:abstractNumId"))
    override = num_2.find(qn("w:lvlOverride")).find(qn("w:startOverride"))
    assert override.get(qn("w:val")) == "5"
    bullet_abstract, _ = _abstract_for(doc, bullets)
    assert _level(bullet_abstract, 0) == ("bullet", "•")
    assert _level(bullet_abstract, 1) == ("bullet", "◦")
    assert _num_pr(doc.paragraphs[-1]) == (bullets, 1)


def test_numbering_definitions_precede_num_entries():
    doc = build_document({"blocks": [{"type": "list", "items": ["a"]}, {"type": "list", "ordered": True, "items": ["b"]}]})
    tags = [child.tag for child in doc.part.numbering_part.element]
    last_abstract = max(i for i, t in enumerate(tags) if t == qn("w:abstractNum"))
    first_num = min(i for i, t in enumerate(tags) if t == qn("w:num"))
    assert last_abstract < first_num


def test_table_has_a_styled_repeating_header_and_padded_rows(tmp_path):
    doc = _render(
        {
            "blocks": [
                {
                    "type": "table",
                    "headers": ["Service", "Fee (EUR)"],
                    "rows": [
                        ["Setup", 1000],
                        [{"runs": [{"text": "Support", "italic": True}]}, {"text": "200", "align": "right"}],
                        ["Only one cell"],
                    ],
                    "widths": [3, 1],
                }
            ]
        },
        tmp_path,
    )
    table = doc.tables[0]
    assert table.style.name == "Table Grid"
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["Service", "Fee (EUR)"],
        ["Setup", "1000"],
        ["Support", "200"],
        ["Only one cell", ""],
    ]
    header = table.rows[0]
    assert header._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(run.bold for cell in header.cells for run in cell.paragraphs[0].runs)
    assert header.cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "D9D9D9"
    assert table.rows[2].cells[0].paragraphs[0].runs[0].italic is True
    assert table.rows[2].cells[1].paragraphs[0].alignment == 2  # RIGHT
    widths = [table.rows[1].cells[i].width for i in range(2)]
    assert widths[0] == pytest.approx(3 * widths[1], rel=0.01)


def test_table_without_header_and_named_style(tmp_path):
    doc = _render({"blocks": [{"type": "table", "rows": [["a", "b"]], "style": "Light Grid Accent 1"}]}, tmp_path)
    table = doc.tables[0]
    assert table.style.name == "Light Grid Accent 1"
    assert table.rows[0]._tr.trPr is None or table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is None


def test_quote_code_page_break_and_named_paragraph_style(tmp_path):
    doc = _render(
        {
            "blocks": [
                {"type": "quote", "text": "Pacta sunt servanda"},
                {"type": "page_break"},
                {"type": "code", "text": "line 1\nline 2"},
                {"type": "paragraph", "text": "Body", "style": "body text"},
                {"type": "paragraph", "text": "Fallback", "style": "No Such Style"},
            ]
        },
        tmp_path,
    )
    paragraphs = doc.paragraphs
    assert paragraphs[0].style.name == "Quote"
    assert any(br.get(qn("w:type")) == "page" for br in doc.element.body.iter(qn("w:br")))
    code = next(p for p in paragraphs if p.text.startswith("line 1"))
    assert code.runs[0].font.name == "Courier New"
    assert code._p.findall(".//" + qn("w:br"))  # the newline became a line break
    assert next(p for p in paragraphs if p.text == "Body").style.name == "Body Text"
    assert next(p for p in paragraphs if p.text == "Fallback").style.name == "Normal"


def test_the_sandbox_program_renders_what_the_module_renders(tmp_path):
    spec = {
        "title": "T",
        "blocks": [
            {"type": "list", "ordered": True, "items": ["a", {"text": "b", "items": ["c"]}]},
            {"type": "table", "headers": ["h"], "rows": [["x"]], "widths": [1]},
            {"type": "paragraph", "runs": [{"text": "l", "link": "https://e.com", "bold": True}]},
        ],
    }
    direct = _render(spec, tmp_path)
    program = _render_in_sandbox_program(spec, tmp_path)
    assert program.element.body.xml == direct.element.body.xml


def test_control_characters_are_dropped_not_fatal(tmp_path):
    doc = _render({"title": "a\x00b", "blocks": [{"type": "paragraph", "text": "c\x0bd\x1fe"}]}, tmp_path)
    assert [p.text for p in doc.paragraphs] == ["ab", "cde"]
    assert clean_text(None) == ""


def test_spec_text_is_data_never_code(capfd, tmp_path):
    sentinel = tmp_path / "pwned.txt"
    payload = f"'''; __import__('os').system('echo PWNED > {sentinel}'); print('SHOULD_NOT_PRINT'); '''"
    spec = {
        "title": payload,
        "blocks": [
            {"type": "paragraph", "runs": [payload, {"text": payload, "bold": True}]},
            {"type": "table", "headers": [payload], "rows": [[payload]]},
            {"type": "list", "ordered": True, "items": [payload]},
        ],
    }
    doc = _render_in_sandbox_program(spec, tmp_path)
    captured = capfd.readouterr()
    assert "SHOULD_NOT_PRINT" not in captured.out
    assert not sentinel.exists()
    assert doc.paragraphs[0].text == payload


def test_renderer_leaves_the_shared_kernel_namespace_alone(tmp_path):
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps({"blocks": [{"type": "paragraph", "text": "x"}]}), encoding="utf-8")
    namespace: dict = {}
    program = _RENDERERS["document"].format(spec_path=str(spec_path), out_path=str(tmp_path / "o.docx"))
    exec(compile(program, "<r>", "exec"), namespace, namespace)  # noqa: S102
    assert [k for k in namespace if not k.startswith("__")] == []


# ---------------------------------------------------------------------------
# Schema and editing
# ---------------------------------------------------------------------------


def test_schema_accepts_the_rich_spec_and_the_legacy_one():
    tool = _tool()
    rich = {
        "title": "T",
        "blocks": [
            {"type": "heading", "text": "H", "level": 2},
            {"type": "paragraph", "runs": ["a", {"text": "b", "bold": True, "link": "https://e.com"}]},
            {"type": "list", "ordered": True, "start": 3, "items": ["a", {"text": "b", "items": ["c"]}]},
            {"type": "table", "headers": ["h"], "rows": [[1, None, True, {"text": "x", "align": "center"}]]},
            {"type": "quote", "runs": ["q"]},
            {"type": "code", "text": "c"},
            {"type": "page_break"},
        ],
    }
    assert tool._validate("document", rich)[1] is None
    assert tool._validate("document", {"sections": [{"heading": "h", "paragraphs": ["p"]}]})[1] is None


@pytest.mark.parametrize(
    "spec, message",
    [
        ({"title": "only"}, "one of 'sections' or 'blocks' is required"),
        ({"blocks": [{"type": "memo"}]}, "unknown block type 'memo'"),
        ({"blocks": [{"type": "table", "headers": ["a"]}]}, "'rows' is a required property (at blocks/0)"),
        ({"blocks": [{"type": "paragraph", "text": "x", "bold": True}]}, "'bold' was unexpected) (at blocks/0)"),
        ({"blocks": [{"type": "heading", "text": "x", "level": 7}]}, "(at blocks/0/level)"),
        ({"blocks": [{"type": "paragraph", "runs": [{"text": "x", "link": "javascript:x"}]}]}, "blocks/0/runs/0/link"),
        ({"blocks": [{"type": "table", "rows": [["a"]], "widths": [0]}]}, "blocks/0/widths/0"),
        ({"sections": [{"paragraphs": [{"text": "rich"}]}]}, "sections/0/paragraphs/0"),
    ],
)
def test_schema_errors_name_the_problem_and_where_it_is(spec, message):
    error = _tool()._validate("document", spec)[1]
    assert error is not None
    assert error["error"].startswith("invalid document spec: ")
    assert message in error["error"]


def test_spec_append_adds_blocks_to_a_legacy_document(tmp_path):
    legacy = {"sections": [{"heading": "Terms", "paragraphs": ["As agreed."]}]}
    appended = _apply_spec_append(legacy, {"blocks": [{"type": "table", "rows": [["a", "b"]]}]})["spec"]
    assert _tool()._validate("document", appended)[1] is None
    doc = _render(appended, tmp_path)
    assert [p.text for p in doc.paragraphs] == ["Terms", "As agreed."]
    assert len(doc.tables) == 1


def test_edit_appends_blocks_through_the_tool(monkeypatch):
    tool = _tool()
    monkeypatch.setattr(
        tool,
        "_load_current",
        lambda _id: {"artifact_id": "a-1", "kind": "document", "spec": {"blocks": [{"type": "page_break"}]}},
    )
    seen = {}

    def _reversion(artifact_id, kind, spec, action, title=None):
        seen.update(spec=spec, action=action)
        return {"status": "ok"}

    monkeypatch.setattr(tool, "_reversion", _reversion)
    result = tool.execute_action(
        "edit_artifact", id="A1", spec_append={"blocks": [{"type": "paragraph", "text": "More"}]}
    )
    assert result == {"status": "ok"}
    assert seen["spec"]["blocks"] == [{"type": "page_break"}, {"type": "paragraph", "text": "More"}]


def test_empty_tables_and_bad_widths_are_skipped_quietly():
    doc = build_document(
        {
            "blocks": [
                {"type": "table", "rows": []},
                {"type": "table", "rows": [["a", "b"]], "widths": ["wide", 1]},
                {"type": "table", "rows": [["a", "b"]], "widths": [1]},
                {"type": "memo", "text": "unknown"},
            ]
        }
    )
    assert len(doc.tables) == 2
    # Widths that can't be used leave python-docx's even split in place.
    for table in doc.tables:
        first, second = table.rows[0].cells
        assert first.width == second.width


def test_an_item_of_the_wrong_type_still_names_where_it_is():
    error = _tool()._validate("document", {"blocks": [{"type": "list", "items": [5]}]})[1]
    assert "(at blocks/0/items/0)" in error["error"]
