"""Smoke tests for the synthetic many-attachments fixtures.

The default run generates every scenario at the ``small`` profile (a few
seconds) and checks the files open, hit their token targets, dedupe by hash,
and that scans have no text layer. The production-sized run is opt-in::

    MANY_ATTACHMENTS_FULL=1 KMP_DUPLICATE_LIB_OK=TRUE python -m pytest \
        tests/test_many_attachments_fixtures.py -m slow
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
import zipfile
from pathlib import Path
from typing import Any, Dict

import pytest

from tests.fixtures.many_attachments import corpus
from tests.fixtures.many_attachments import generate as gen
from tests.fixtures.many_attachments import v1_requests as v1

DISPOSITIONS = {"inline", "back_filled", "partial", "image", "tool", "sandbox", "earlier", "not_included",
                "unreadable", "skipped", "duplicate"}


def _text_of(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if suffix in (".html", ".xml"):
        return re.sub(r"<[^>]+>", " ", path.read_text(encoding="utf-8")) if suffix == ".html" else path.read_text()
    if suffix in (".epub", ".zip"):
        with zipfile.ZipFile(path) as zf:
            names = [n for n in zf.namelist() if n.endswith((".xhtml", ".tsx", ".md", ".json", ".js", ".css"))]
            return " ".join(re.sub(r"<[^>]+>", " ", zf.read(n).decode("utf-8", "ignore")) for n in names)
    return path.read_text(encoding="utf-8")


def _within(actual: int, target: int, tolerance: float) -> bool:
    return abs(actual - target) <= max(tolerance * target, 150)


@pytest.fixture(scope="module")
def small_set(tmp_path_factory: pytest.TempPathFactory) -> Dict[str, Any]:
    out = tmp_path_factory.mktemp("many_attachments")
    return {"out": out, "manifests": gen.generate(out)}


# ---------------------------------------------------------------------------
# Scenario catalogue (no files needed)
# ---------------------------------------------------------------------------


class TestScenarioCatalogue:
    def test_every_scenario_resolves(self) -> None:
        data = gen.load_scenarios()
        ids = gen.scenario_ids()
        assert {f"RC-{i:02d}" for i in range(1, 13)} | {f"V1-0{i}" for i in range(1, 6)} == set(ids)
        for sid in ids:
            scenario = gen.get_scenario(sid)
            assert scenario["model"] in data["models"]
            specs = gen.file_specs(scenario)
            assert specs, sid
            assert all(s.kind in gen.WRITERS for s in specs), sid
            keys = {s.key for s in specs}
            for conv in gen.conversations(scenario):
                for turn in conv:
                    assert set(turn.attach) <= keys, (sid, turn.index)

    def test_every_expectation_rule_applies_somewhere(self) -> None:
        for sid in gen.scenario_ids():
            scenario = gen.get_scenario(sid)
            convs = gen.conversations(scenario)
            matched = set()
            for caps in gen.capability_matrix(scenario):
                for conv in convs:
                    for turn in conv:
                        exp = gen.expectations_for(scenario, turn, caps)
                        for allowed in exp["files"].values():
                            assert set(allowed) <= DISPOSITIONS, (sid, allowed)
                        for i, rule in enumerate(scenario.get("expect", [])):
                            single = {**scenario, "expect": [rule]}
                            got = gen.expectations_for(single, turn, caps)
                            if got["files"] or got["props"]:
                                matched.add(i)
            assert matched == set(range(len(scenario.get("expect", [])))), (sid, matched)

    def test_turn_shapes_match_real_cases(self) -> None:
        rc01 = gen.conversations(gen.get_scenario("RC-01"))[0]
        assert [len(t.attach) for t in rc01] == [1, 5, 15, 18, 5, 2, 8, 9, 0]
        assert rc01[4].uploading == 2
        rc04 = gen.conversations(gen.get_scenario("RC-04"))[0]
        assert len(rc04[1].attach) == 46
        rc08 = gen.conversations(gen.get_scenario("RC-08"))[0]
        assert len(rc08) == 126 and rc08[0].attach == ["quiz[0]"] and rc08[125].attach == []

    def test_expectations_resolve_new_and_earlier(self) -> None:
        scenario = gen.get_scenario("RC-01")
        conv = gen.conversations(scenario)[0]
        caps = {"tool_calling": True, "sandbox": False, "vision": True}
        last = gen.expectations_for(scenario, conv[8], caps)
        assert len(last["files"]) == 60 and set(map(tuple, last["files"].values())) == {("earlier",)}
        assert last["props"]["mentions_tools"] is True
        dupes = gen.expectations_for(scenario, conv[6], caps)
        assert dupes["props"]["duplicates"] == 3
        assert dupes["files"]["invoice_dupes[0]"] == ["duplicate"]
        assert gen.disposition_matches(["inline"], "back_filled")
        assert not gen.disposition_matches(["back_filled"], "inline")
        assert set(gen.PLANNER_STATUS) == DISPOSITIONS
        assert gen.planner_statuses(["image", "partial"]) == {"inline", "partial"}

    def test_attachment_records_carry_full_sizes_and_shared_hashes(self) -> None:
        rows = gen.attachment_records(gen.get_scenario("RC-04"))
        assert len(rows) == 46
        by_key = {r["key"]: r for r in rows}
        assert by_key["exam_dupes[0]"]["content_hash"] == by_key["exams[4]"]["content_hash"]
        assert len({r["content_hash"] for r in rows}) == 44
        reports = [r["token_count"] for r in rows if r["group"] == "reports"]
        assert min(reports) >= 128000 and max(reports) <= 454000
        assert sum(r["token_count"] for r in rows if r["group"] == "exams") < 130000

    def test_capability_matrix_covers_variants(self) -> None:
        caps = gen.capability_matrix(gen.get_scenario("RC-01"))
        assert {(c["tool_calling"], c["sandbox"], c["vision"]) for c in caps} == {
            (True, True, True), (True, True, False), (True, False, True), (True, False, False),
            (False, False, True), (False, False, False)}
        v1_caps = gen.capability_matrix(gen.get_scenario("V1-01"))
        assert {c["payload"] for c in v1_caps} == {"text", "file_parts"}


# ---------------------------------------------------------------------------
# Generated files (small profile)
# ---------------------------------------------------------------------------


class TestGeneratedFiles:
    def test_files_exist_and_open(self, small_set: Dict[str, Any]) -> None:
        from openpyxl import load_workbook
        from PIL import Image
        from pptx import Presentation
        from pypdf import PdfReader

        out = small_set["out"]
        for sid, manifest in small_set["manifests"].items():
            for f in manifest["files"]:
                path = out / sid / f["path"]
                assert path.is_file(), f["path"]
                assert hashlib.sha256(path.read_bytes()).hexdigest() == f["sha256"]
                suffix = path.suffix.lower()
                if f["kind"] == "bad_file":
                    with pytest.raises(Exception):
                        Image.open(path).load()
                elif suffix == ".pdf":
                    assert len(PdfReader(path).pages) == f["pages"] >= 1
                elif suffix in (".png", ".jpg", ".webp"):
                    with Image.open(path) as img:
                        img.load()
                        assert img.format == {".png": "PNG", ".jpg": "JPEG", ".webp": "WEBP"}[suffix]
                elif suffix in (".xlsx", ".xlsm"):
                    wb = load_workbook(path, read_only=True)
                    assert wb.sheetnames == f["sheets"]
                elif suffix == ".csv":
                    with open(path, newline="", encoding="utf-8") as fh:
                        assert sum(1 for _ in csv.reader(fh)) == f["rows"] + 1
                elif suffix == ".pptx":
                    assert len(Presentation(str(path)).slides) == f["slides"]
                elif suffix in (".zip", ".epub"):
                    with zipfile.ZipFile(path) as zf:
                        assert zf.testzip() is None

    def test_text_files_hit_token_targets(self, small_set: Dict[str, Any]) -> None:
        tolerance = gen.load_scenarios()["token_tolerance"]
        misses = []
        for sid, manifest in small_set["manifests"].items():
            for f in manifest["files"]:
                target = f.get("target_tokens")
                if not target or not f["has_text_layer"]:
                    continue
                actual = corpus.count_tokens(_text_of(small_set["out"] / sid / f["path"]))
                if not _within(actual, target, tolerance):
                    misses.append((sid, f["key"], target, actual))
        assert not misses

    def test_spreadsheets_hit_row_targets(self, small_set: Dict[str, Any]) -> None:
        rc03 = {f["key"]: f for f in small_set["manifests"]["RC-03"]["files"]}
        assert rc03["workbooks[0]"]["rows"] == 300 and rc03["workbooks[0]"]["full_rows"] == 20000
        assert rc03["tables[1]"]["rows"] == 300 and rc03["tables[2]"]["full_rows"] == 200000
        assert rc03["workbooks[0]"]["sheets"] == ["raw_data", "summary", "lookup_1"]

    def test_duplicates_hash_equal_their_sources(self, small_set: Dict[str, Any]) -> None:
        count = 0
        for manifest in small_set["manifests"].values():
            by_key = {f["key"]: f for f in manifest["files"]}
            for f in manifest["files"]:
                if f.get("duplicate_of"):
                    assert f["sha256"] == by_key[f["duplicate_of"]]["sha256"]
                    count += 1
        assert count == 5  # RC-01: 3, RC-04: 2

    def test_scans_have_no_text_layer(self, small_set: Dict[str, Any]) -> None:
        from pypdf import PdfReader

        scans = 0
        for sid, manifest in small_set["manifests"].items():
            for f in manifest["files"]:
                if f["kind"] in gen.SCAN_KINDS:
                    reader = PdfReader(small_set["out"] / sid / f["path"])
                    assert all(not (p.extract_text() or "").strip() for p in reader.pages), f["path"]
                    assert all(p.images for p in reader.pages)
                    scans += 1
        assert scans == 4 + 2 + 3 + 2 + 2  # RC-01, RC-10, RC-12, V1-01, V1-04

    def test_zip_archives_have_expected_members(self, small_set: Dict[str, Any]) -> None:
        sdir = small_set["out"] / "RC-11"
        with zipfile.ZipFile(sdir / "zip_mixed" / "handover_bundle.zip") as zf:
            names = set(zf.namelist())
            assert {"README.txt", "notes/meeting.md", "data/members.csv", "images/diagram.png",
                    "docs/summary.pdf", "notes/blob.bin", "inner.zip"} == names
            with zipfile.ZipFile(__import__("io").BytesIO(zf.read("inner.zip"))) as inner:
                assert set(inner.namelist()) == {"inner_a.txt", "inner_b.md"}
        with zipfile.ZipFile(sdir / "zip_app" / "storefront_app.zip") as zf:
            assert len(zf.namelist()) == 40 and "public/logo.png" in zf.namelist()
        with zipfile.ZipFile(sdir / "zip_docs" / "product_docs.zip") as zf:
            assert len(zf.namelist()) == 25

    def test_hindi_and_slovenian_text(self, small_set: Dict[str, Any]) -> None:
        rc06 = small_set["manifests"]["RC-06"]
        text = _text_of(small_set["out"] / "RC-06" / rc06["files"][0]["path"])
        if rc06["devanagari_font"]:
            assert re.search(r"[ऀ-ॿ]{3,}", text)
        else:
            assert "Niyam" in text
        odlok = _text_of(small_set["out"] / "V1-01" / "odlok" / "Odlok_OPN_Primerjevo.pdf")
        assert "člen" in odlok and "FIKTIVNO" in odlok

    def test_identifiers_are_placeholders(self, small_set: Dict[str, Any]) -> None:
        for sid in small_set["manifests"]:
            for tp in (small_set["out"] / sid / "_text").rglob("*.txt"):
                text = tp.read_text(encoding="utf-8")
                assert not re.search(r"[\w.]+@[\w-]+\.\w+", text), tp
                for long_id in re.findall(r"\b\d{10,13}\b", text):
                    assert set(long_id) == {"0"}, (tp, long_id)

    def test_output_is_deterministic(self, tmp_path: Path, small_set: Dict[str, Any]) -> None:
        for sid in ("RC-03", "RC-11"):
            again = gen.generate_scenario(sid, tmp_path, "small")
            first = {f["key"]: f["sha256"] for f in small_set["manifests"][sid]["files"]}
            assert {f["key"]: f["sha256"] for f in again["files"]} == first


# ---------------------------------------------------------------------------
# /v1 request shapes
# ---------------------------------------------------------------------------


class TestV1Requests:
    def test_cumulative_text_reattachment(self, small_set: Dict[str, Any]) -> None:
        steps = v1.build_requests("V1-01", small_set["out"], payload="text")
        assert [s["step"] for s in steps] == ["user"] * 5
        contents = [s["body"]["messages"][-1]["content"] for s in steps]
        assert [c.count("--- Začetek besedila:") for c in contents] == [2, 3, 6, 8, 16]
        assert "[Priloga izpisek_skeniran_1.pdf: skeniran dokument (2 strani)" in contents[2]
        assert [t["function"]["name"] for t in steps[0]["body"]["tools"]] == [
            "search", "get_by_id", "create_artefact", "str_replace_editor"]
        assert steps[0]["body"]["messages"][0]["role"] == "system" and steps[0]["body"]["stream"] is True

    def test_file_parts_reattachment(self, small_set: Dict[str, Any]) -> None:
        steps = v1.build_requests("V1-01", small_set["out"], payload="file_parts")
        parts = [s["body"]["messages"][-1]["content"] for s in steps]
        assert [sum(p["type"] == "file" for p in c) for c in parts] == [2, 3, 6, 8, 16]
        assert parts[0][1]["file"]["file_data"].startswith("data:application/pdf;base64,")
        assert any(p["file"]["filename"].endswith(".docx") for p in parts[1] if p["type"] == "file")

    def test_upper_case_pdf_parts_and_tool_replay(self, small_set: Dict[str, Any]) -> None:
        steps = v1.build_requests("V1-03", small_set["out"], payload="file_parts")
        assert [s["step"] for s in steps] == ["user", "tool_round", "tool_round", "tool_round"]
        first = steps[0]["body"]["messages"][-1]["content"]
        names = [p["file"]["filename"] for p in first if p["type"] == "file"]
        assert names == ["PRILOGA_1.PDF", "PRILOGA_2.PDF", "PRILOGA_3.PDF", "PRILOGA_4.PDF"]
        assert all(p["file"]["file_data"].startswith("data:application/pdf;base64,") for p in first[1:])
        last = steps[-1]["body"]["messages"]
        assert last[1]["content"] == first  # continuation replays the file parts
        assert [m["role"] for m in last[2:]] == ["assistant", "tool"] * 3

    def test_image_parts_and_forced_final(self, small_set: Dict[str, Any]) -> None:
        steps = v1.build_requests("V1-02", small_set["out"], payload="file_parts")
        content = steps[0]["body"]["messages"][-1]["content"]
        assert sum(p["type"] == "image_url" for p in content) == 10
        assert sum(p["type"] == "file" for p in content) == 2
        loop = v1.build_requests("V1-04", small_set["out"], payload="file_parts")
        assert [s["step"] for s in loop] == ["user"] + ["tool_round"] * 4 + ["forced_final"]
        assert "tools" not in loop[-1]["body"]
        assert sum(p["type"] == "file" for p in loop[-1]["body"]["messages"][-1]["content"]) == 8

    def test_stateless_classifier_batch(self, small_set: Dict[str, Any]) -> None:
        for payload in ("text", "file_parts"):
            steps = v1.build_requests("V1-05", small_set["out"], payload=payload)
            assert len(steps) == 20 and all(s["stateless"] for s in steps)
            for s in steps:
                assert [m["role"] for m in s["body"]["messages"]] == ["user"]
                assert "tools" not in s["body"] and s["body"]["stream"] is False
        texts = [s["body"]["messages"][0]["content"] for s in v1.build_requests("V1-05", small_set["out"])]
        assert len(set(texts)) == 20 and all("PRODAJNA POGODBA" in t for t in texts)


# ---------------------------------------------------------------------------
# Production-sized files (opt-in)
# ---------------------------------------------------------------------------


@pytest.mark.slow
@pytest.mark.skipif(not os.environ.get("MANY_ATTACHMENTS_FULL"), reason="set MANY_ATTACHMENTS_FULL=1")
def test_full_profile_long_documents(tmp_path: Path) -> None:
    tolerance = gen.load_scenarios()["token_tolerance"]
    picks = {"RC-06": ["rulebook_a"], "V1-02": ["odlok_big"], "RC-05": ["papers"], "RC-12": ["scanned_books"]}
    for sid, groups in picks.items():
        manifest = gen.generate_scenario(sid, tmp_path, "full", groups=groups)
        for f in manifest["files"]:
            if f["kind"] in gen.SCAN_KINDS:
                assert f["pages"] == 400
                continue
            actual = corpus.count_tokens(_text_of(tmp_path / sid / f["path"]))
            assert _within(actual, f["target_tokens"], tolerance), (sid, f["key"], f["target_tokens"], actual)
    thesis = gen.load_manifest(tmp_path, "RC-05")["files"][8]
    assert thesis["thesis"] and thesis["pages"] >= 120
