import pytest
from pathlib import Path
from unittest.mock import patch

from docsgpt.parser.file.pptx_parser import PPTXParser


def test_pptx_init_parser():
    parser = PPTXParser()
    assert isinstance(parser._init_parser(), dict)
    assert not parser.parser_config_set
    parser.init_parser()
    assert parser.parser_config_set


def _fake_presentation_with(slides_shapes_texts):
    class Shape:
        def __init__(self, text=None):
            if text is not None:
                self.text = text
    class Slide:
        def __init__(self, texts):
            self.shapes = [Shape(t) for t in texts]
    class Pres:
        def __init__(self, _file):
            self.slides = [Slide(texts) for texts in slides_shapes_texts]
    return Pres


def test_pptx_parser_concat_true():
    # The shapes are separate blocks of text, so they are joined with a newline
    # rather than run together.
    slides = [["Hello ", "World"], ["Slide2"]]
    FakePres = _fake_presentation_with(slides)
    import sys
    import types
    fake_pptx = types.ModuleType("pptx")
    fake_pptx.Presentation = FakePres
    parser = PPTXParser()
    with patch.dict(sys.modules, {"pptx": fake_pptx}):
        result = parser.parse_file(Path("deck.pptx"))
    assert result == "Hello\nWorld\nSlide2"


def test_pptx_parser_list_mode():
    slides = [[" A ", "B"], [" C "]]
    FakePres = _fake_presentation_with(slides)
    import sys
    import types
    fake_pptx = types.ModuleType("pptx")
    fake_pptx.Presentation = FakePres
    parser = PPTXParser()
    parser._concat_slides = False
    with patch.dict(sys.modules, {"pptx": fake_pptx}):
        result = parser.parse_file(Path("deck.pptx"))
    assert result == ["A\nB", "C"]


def test_pptx_parser_import_error():
    parser = PPTXParser()
    import sys
    with patch.dict(sys.modules, {"pptx": None}):
        with pytest.raises(ImportError, match="pptx module is required to read .PPTX files"):
            parser.parse_file(Path("missing.pptx"))


# --- groups, tables and the shape separator ------------------------------------


def _deck_path(tmp_path, build):
    """Save a deck built by `build(slide)` on one blank slide and return its path."""
    from pptx import Presentation

    presentation = Presentation()
    build(presentation.slides.add_slide(presentation.slide_layouts[6]))
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def test_shapes_do_not_run_into_each_other(tmp_path):
    """Concatenating without a separator glues the last word of one shape to the
    first word of the next, which is then embedded as a compound word."""
    pytest.importorskip("pptx")
    from pptx.util import Inches

    def build(slide):
        first = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(4), Inches(0.5))
        first.text_frame.text = "FIRST BOX"
        second = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(4), Inches(0.5))
        second.text_frame.text = "SECOND BOX"

    text = PPTXParser().parse_file(_deck_path(tmp_path, build))

    assert "BOXSECOND" not in text
    assert "FIRST BOX" in text
    assert "SECOND BOX" in text


def test_table_cells_are_extracted(tmp_path):
    """A table lives on a graphic frame, which carries no `text` of its own."""
    pytest.importorskip("pptx")
    from pptx.util import Inches

    def build(slide):
        table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(0.5), Inches(6), Inches(1.2)).table
        table.cell(0, 0).text = "HEADER A"
        table.cell(0, 1).text = "HEADER B"
        table.cell(1, 0).text = "CELL A"
        table.cell(1, 1).text = "CELL B"

    text = PPTXParser().parse_file(_deck_path(tmp_path, build))

    assert "| HEADER A | HEADER B |" in text
    assert "| CELL A | CELL B |" in text


def test_a_cell_cannot_break_the_table_row(tmp_path):
    """A pipe would add a column and a line break would end the row early."""
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(0.5), Inches(6), Inches(1.2)).table
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Notes"
    table.cell(1, 0).text = "a|b"
    cell = table.cell(1, 1)
    cell.text = "first line"
    cell.text_frame.add_paragraph().text = "second line"

    path = tmp_path / "cells.pptx"
    presentation.save(str(path))

    text = PPTXParser().parse_file(path)
    rows = [row for row in text.split("\n") if row.startswith("|")]

    assert rows == [
        "| Step | Notes |",
        "| --- | --- |",
        "| a\\|b | first line second line |",
    ]


def test_text_inside_a_group_is_extracted(tmp_path):
    """A group carries no text of its own; its children do."""
    pytest.importorskip("pptx")
    from pptx.util import Inches

    def build(slide):
        left = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.5))
        left.text_frame.text = "GROUPED ONE"
        right = slide.shapes.add_textbox(Inches(3.0), Inches(0.5), Inches(2), Inches(0.5))
        right.text_frame.text = "GROUPED TWO"
        slide.shapes.add_group_shape([left, right])

    text = PPTXParser().parse_file(_deck_path(tmp_path, build))

    assert "GROUPED ONE" in text
    assert "GROUPED TWO" in text


def test_empty_slide_produces_an_empty_string(tmp_path):
    pytest.importorskip("pptx")

    assert PPTXParser().parse_file(_deck_path(tmp_path, lambda slide: None)) == ""
