#!/usr/bin/env python3
"""Smoke-test a sandbox image from the inside: every manifest item, plus real conversions.

Run it in the image the way kernels run, so the environment checks see what
model-written code sees::

    # self-hosted runner image (read-only root, like compose and k8s):
    docker run --rm --read-only --tmpfs /tmp IMAGE \\
        /opt/docsgpt/kernel-env.sh python /opt/docsgpt/smoke_test.py
    # Daytona snapshot: scripts/build_daytona_snapshot.py --smoke

It reads manifest.json (written from docsgpt/sandbox/manifest.py, next to this
script in the image) and checks that every Python package imports, every
command is on PATH, every font file exists and the image environment is set.
Then it converts real files: docx and pptx to PDF with office-convert, HTML to
PDF and PNG with headless Chromium, OCR of a rendered image, pdftotext and
pdfplumber on a generated PDF, Node and npm, an animated GIF and WebP through
imageio, and an H.264 MP4 through the ffmpeg command, checked with ffprobe.
``--pip`` also pip-installs a small package as the sandbox user and imports it
in the running interpreter (needs network).

Each check prints PASS or FAIL; the exit code is 1 if any failed. Only the
stdlib and the image's own packages are used.
"""

from __future__ import annotations

import argparse
import importlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

MARKER = "DocsGPT sandbox smoke 4217"
_DEFAULT_MANIFEST = Path(__file__).resolve().with_name("manifest.json")


class SmokeFailure(Exception):
    """A check found the image wrong."""


def _run(cmd: List[str], timeout: float = 120, **kwargs) -> subprocess.CompletedProcess:
    """Run a command, raising SmokeFailure with its output when it fails."""
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, **kwargs)
    if proc.returncode != 0:
        raise SmokeFailure(f"{' '.join(cmd)} exited {proc.returncode}: {(proc.stderr or proc.stdout).strip()[-600:]}")
    return proc


def _pdf_text(path: Path) -> str:
    """Return the text of a PDF via pypdf."""
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _expect(condition: bool, message: str) -> None:
    if not condition:
        raise SmokeFailure(message)


# -- Manifest checks -----------------------------------------------------------


def check_imports(manifest: Dict) -> str:
    missing = []
    for pkg in manifest["pip"]:
        try:
            importlib.import_module(pkg["import"])
        except Exception as exc:  # noqa: BLE001 - report every broken import
            missing.append(f"{pkg['import']} ({type(exc).__name__}: {exc})")
    _expect(not missing, "cannot import: " + "; ".join(missing))
    return f"{len(manifest['pip'])} packages"


def check_binaries(manifest: Dict) -> str:
    on_path = [b["name"] for b in manifest["binaries"] if b["on_path"]]
    missing = [name for name in on_path if shutil.which(name) is None]
    _expect(not missing, "not on PATH: " + ", ".join(missing))
    return f"{len(on_path)} commands"


def check_fonts(manifest: Dict) -> str:
    missing = [f["path"] for f in manifest["fonts"] if not Path(f["path"]).is_file()]
    _expect(not missing, "missing font files: " + ", ".join(missing))
    families = _run(["fc-list", ":", "family"]).stdout
    for family in ("DejaVu Sans", "Liberation Sans", "Carlito", "Noto Sans Arabic", "Noto Sans CJK"):
        _expect(family in families, f"fontconfig does not list {family}")
    # The liberation2 compatibility link must not make fontconfig list the fonts twice.
    regular = [line for line in _run(["fc-list"]).stdout.splitlines() if "LiberationSans-Regular" in line]
    _expect(len(regular) == 1, f"Liberation Sans listed {len(regular)} times")
    # The fonts the code executor offers for reportlab must load there.
    from reportlab.pdfbase.ttfonts import TTFont

    for number, font in enumerate(f for f in manifest["fonts"] if f.get("reportlab")):
        TTFont(f"smoke{number}", font["path"])
    return f"{len(manifest['fonts'])} font files"


def check_env(manifest: Dict) -> str:
    wrong = {k: os.environ.get(k) for k, v in manifest["env"].items() if os.environ.get(k) != v}
    _expect(not wrong, f"env not set as in the manifest: {wrong}")
    return ", ".join(f"{k}={v}" for k, v in manifest["env"].items())


def check_writable_home(manifest: Dict) -> str:
    home = Path.home()
    probe = home / ".cache" / "docsgpt-smoke"
    probe.parent.mkdir(parents=True, exist_ok=True)
    probe.write_text("ok")
    probe.unlink()
    return str(home)


# -- Conversions ---------------------------------------------------------------


def check_office_docx_to_pdf(work: Path) -> str:
    import docx

    doc = docx.Document()
    doc.add_heading("Smoke", 1)
    doc.add_paragraph(MARKER)
    src = work / "report.docx"
    doc.save(src)
    out = Path(_run(["office-convert", str(src), "--to", "pdf", "--outdir", str(work)]).stdout.strip())
    _expect(out == work / "report.pdf", f"unexpected output path {out}")
    _expect(MARKER in _pdf_text(out), "converted PDF lacks the marker text")
    return f"{out.name} {out.stat().st_size} bytes"


def check_office_pptx_to_pdf(work: Path) -> str:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Smoke"
    slide.placeholders[1].text = MARKER
    src = work / "deck.pptx"
    prs.save(src)
    # Two conversions back to back: a shared LibreOffice profile would make the second silently write nothing.
    _run(["office-convert", str(src), "--to", "pdf", "--outdir", str(work / "a")])
    out = Path(_run(["office-convert", str(src), "--to", "pdf", "--outdir", str(work / "b")]).stdout.strip())
    _expect(MARKER in _pdf_text(out), "converted PDF lacks the marker text")
    return f"{out.name} {out.stat().st_size} bytes"


def _write_html(work: Path) -> Path:
    page = work / "page.html"
    page.write_text(
        "<!doctype html><html><head><meta charset='utf-8'><style>body{font-family:sans-serif}</style></head>"
        f"<body><h1>{MARKER}</h1><p lang='zh'>中文 日本語</p><p lang='ar' dir='rtl'>مرحبا</p><p lang='hi'>नमस्ते</p>"
        "<script>document.body.insertAdjacentHTML('beforeend','<p>script ran</p>')</script>"
        "</body></html>",
        encoding="utf-8",
    )
    return page


def check_html_to_pdf(work: Path) -> str:
    page = _write_html(work)
    out = work / "page.pdf"
    _run(["html-to-pdf", str(page), str(out)])
    text = _run(["pdftotext", str(out), "-"]).stdout
    _expect(MARKER in text, "PDF lacks the marker text")
    _expect("script ran" in text, "page script did not run before printing")
    return f"{out.stat().st_size} bytes"


def check_html_screenshot(work: Path) -> str:
    from PIL import Image

    page = _write_html(work)
    out = work / "shot.png"
    _run(["html-screenshot", str(page), str(out), "--width", "640", "--height", "480"])
    with Image.open(out) as img:
        _expect(img.size == (640, 480), f"screenshot is {img.size}, expected (640, 480)")
        _expect(len(img.convert("L").getcolors(1 << 16) or []) > 2, "screenshot is blank")
    return "640x480"


def _render_text_png(work: Path, text: str, font_path: str) -> Path:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1400, 160), "white")
    ImageDraw.Draw(img).text((30, 40), text, fill="black", font=ImageFont.truetype(font_path, 64))
    out = work / "ocr.png"
    img.save(out)
    return out


def check_ocr(work: Path) -> str:
    import pytesseract

    image = _render_text_png(work, MARKER, "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    text = pytesseract.image_to_string(str(image))
    _expect("sandbox smoke 4217" in text, f"tesseract read {text.strip()!r}")
    return f"tesseract {pytesseract.get_tesseract_version()}"


def _reportlab_pdf(work: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    out = work / "generated.pdf"
    c = canvas.Canvas(str(out), pagesize=A4)
    c.drawString(72, 760, MARKER)
    c.save()
    return out


def check_pdf_tools(work: Path) -> str:
    import pdfplumber
    import pypdfium2

    pdf = _reportlab_pdf(work)
    _expect(MARKER in _run(["pdftotext", str(pdf), "-"]).stdout, "pdftotext lost the text")
    with pdfplumber.open(pdf) as doc:
        _expect(MARKER in (doc.pages[0].extract_text() or ""), "pdfplumber lost the text")
        doc.pages[0].to_image(resolution=50).save(work / "plumber.png")
    _run(["pdftoppm", "-png", "-r", "50", "-singlefile", str(pdf), str(work / "page")])
    _expect((work / "page.png").is_file(), "pdftoppm wrote no image")
    pypdfium2.PdfDocument(str(pdf))[0].render(scale=0.5).to_pil().save(work / "pdfium.png")
    return "pdftotext, pdfplumber, pdftoppm, pypdfium2"


def check_node(manifest: Dict) -> str:
    version = _run(["node", "--version"]).stdout.strip()
    _expect(version == f"v{manifest['node']['version']}", f"node is {version}")
    npm = _run(["npm", "--version"]).stdout.strip()
    _run(["npx", "--version"])
    _expect(_run(["node", "-e", "console.log(String(6 * 7))"]).stdout.strip() == "42", "node -e failed")
    return f"node {version}, npm {npm}"


def check_animation(work: Path) -> str:
    import imageio.v3 as iio
    import numpy as np
    from PIL import Image

    frames = np.zeros((6, 48, 64, 3), dtype=np.uint8)
    for i in range(len(frames)):
        frames[i, :, i * 10 : i * 10 + 10] = 255
    results = []
    for name in ("anim.gif", "anim.webp"):
        out = work / name
        iio.imwrite(out, frames, duration=100, loop=0)
        with Image.open(out) as img:
            _expect(getattr(img, "n_frames", 1) == len(frames), f"{name} has {getattr(img, 'n_frames', 1)} frames")
        results.append(f"{name} {out.stat().st_size} bytes")
    return ", ".join(results)


def check_ffmpeg_mp4(work: Path) -> str:
    import numpy as np

    out = work / "clip.mp4"
    frames = np.zeros((10, 240, 320, 3), dtype=np.uint8)
    for i in range(len(frames)):
        frames[i, :, i * 32 : i * 32 + 32] = (255, 128, 0)
    cmd = [
        "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
        "-f", "rawvideo", "-pix_fmt", "rgb24", "-s", "320x240", "-r", "10", "-i", "-",
        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(out),
    ]  # fmt: skip
    proc = subprocess.run(cmd, input=frames.tobytes(), capture_output=True, timeout=120)
    _expect(proc.returncode == 0, f"ffmpeg exited {proc.returncode}: {proc.stderr.decode()[-600:]}")
    probe = json.loads(
        _run(["ffprobe", "-v", "error", "-show_entries", "stream=codec_name,nb_frames", "-of", "json", str(out)]).stdout
    )
    stream = probe["streams"][0]
    _expect(stream["codec_name"] == "h264", f"codec is {stream['codec_name']}")
    _expect(int(stream.get("nb_frames", 0)) == len(frames), f"{stream.get('nb_frames')} frames")
    return f"h264, {len(frames)} frames, {out.stat().st_size} bytes"


def check_pip_user_install(work: Path) -> str:
    """pip install as the sandbox user, then import in this interpreter (what a kernel does)."""
    pkg = "tabulate==0.9.0"
    _run([sys.executable, "-m", "pip", "install", "--no-deps", "--quiet", pkg], timeout=180)
    importlib.invalidate_caches()
    module = importlib.import_module("tabulate")
    return f"tabulate imported from {Path(module.__file__).parent}"


# -- Runner --------------------------------------------------------------------


def _checks(manifest: Dict, work: Path, with_pip: bool) -> List[Tuple[str, Callable[[], str]]]:
    def sub(name: str) -> Path:
        path = work / name
        path.mkdir()
        return path

    checks = [
        ("imports", lambda: check_imports(manifest)),
        ("commands on PATH", lambda: check_binaries(manifest)),
        ("fonts", lambda: check_fonts(manifest)),
        ("environment", lambda: check_env(manifest)),
        ("writable HOME", lambda: check_writable_home(manifest)),
        ("office-convert docx -> pdf", lambda: check_office_docx_to_pdf(sub("docx"))),
        ("office-convert pptx -> pdf (twice)", lambda: check_office_pptx_to_pdf(sub("pptx"))),
        ("html-to-pdf", lambda: check_html_to_pdf(sub("html-pdf"))),
        ("html-screenshot", lambda: check_html_screenshot(sub("html-png"))),
        ("OCR (pytesseract)", lambda: check_ocr(sub("ocr"))),
        ("PDF tools", lambda: check_pdf_tools(sub("pdf"))),
        ("node / npm / npx", lambda: check_node(manifest)),
        ("animated GIF / WebP (imageio)", lambda: check_animation(sub("anim"))),
        ("MP4 (ffmpeg + ffprobe)", lambda: check_ffmpeg_mp4(sub("mp4"))),
    ]
    if with_pip:
        checks.append(("pip install --user + import", lambda: check_pip_user_install(sub("pip"))))
    return checks


def main(argv: Optional[List[str]] = None) -> int:
    """Run every check and print a PASS/FAIL line for each.

    Args:
        argv: Command-line arguments; defaults to ``sys.argv[1:]``.

    Returns:
        0 when every check passed, 1 otherwise.
    """
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--manifest", type=Path, default=_DEFAULT_MANIFEST, help="path to manifest.json")
    parser.add_argument("--pip", action="store_true", help="also pip-install a package and import it (network)")
    args = parser.parse_args(argv)
    manifest = json.loads(args.manifest.read_text())

    failed = 0
    started = time.monotonic()
    with tempfile.TemporaryDirectory(prefix="smoke-") as tmp:
        for name, check in _checks(manifest, Path(tmp), args.pip):
            t0 = time.monotonic()
            try:
                detail = check()
                status = "PASS"
            except SmokeFailure as exc:
                detail, status = str(exc), "FAIL"
            except Exception as exc:  # noqa: BLE001 - a crash in a check is a failed check
                detail = f"{type(exc).__name__}: {exc}\n{traceback.format_exc(limit=3)}"
                status = "FAIL"
            failed += status == "FAIL"
            print(f"{status} {name} ({time.monotonic() - t0:.1f}s): {detail}", flush=True)
    print(f"{'FAILED' if failed else 'OK'}: {failed} failed, {time.monotonic() - started:.0f}s total")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
